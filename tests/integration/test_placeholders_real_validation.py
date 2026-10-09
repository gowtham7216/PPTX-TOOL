"""
Integration and Real-Presentation Validation Tests for Task 4 Option C.
Validates placeholder conversion and design preservation across:
1. Ch06_GLN_Instructor_AU.pptx (50 slides, native placeholders, animations, groups)
2. Sec1-5-TrigFunctions (2)original.pptx (Math diagrams, tables, shapes)
3. Chapter 9.pptx (Chapter 9 deck with numeric badges '03', '04', cards, complex components)
"""

from pathlib import Path
import pytest
from pptx import Presentation
from lxml import etree

from pptx_engine.reader.presentation_reader import PresentationReader
from pptx_engine.writer.presentation_writer import PresentationWriter
from rules.placeholders.classifier import (
    PlaceholderClassifier,
    PlaceholderCategory,
    ConversionAction,
    _safe_xpath,
    NSMAP,
)
from rules.placeholders.transformer import PlaceholderTransformer
from rules.placeholders.rule import PlaceholderRule
from models.enums import RuleBehavior, RemediationStatus


@pytest.fixture
def reader():
    return PresentationReader()


@pytest.fixture
def rule():
    return PlaceholderRule()


# =============================================================================
# 1. Validation Deck: Ch06_GLN_Instructor_AU.pptx
# =============================================================================
def test_real_validation_ch06(tmp_path, reader, rule):
    deck_path = Path("tests/fixtures/Ch06_GLN_Instructor_AU.pptx")
    if not deck_path.exists():
        pytest.skip(f"{deck_path} not found.")

    pres_model = reader.read(deck_path)
    initial_slide_count = len(pres_model.slides)
    assert initial_slide_count == 50

    # Generate comprehensive audit report
    audit_report = rule.generate_presentation_audit_report(pres_model)
    assert audit_report["total_slides_evaluated"] == 50
    assert audit_report["visual_integrity_verified"] is True
    assert audit_report["failed_conversions"] == 0

    # Execute remediation across all slides
    applied_count = 0
    skipped_count = 0
    for slide in pres_model.slides:
        issues = rule.evaluate(slide, pres_model)
        for issue in issues:
            if issue.behavior == RuleBehavior.SAFE_AUTO_REMEDIATE:
                action = rule.remediate_issue(slide, pres_model, issue)
                if action.status == RemediationStatus.APPLIED:
                    applied_count += 1
                elif action.status == RemediationStatus.SKIPPED:
                    skipped_count += 1

    # Save to temp path and verify reload
    out_file = tmp_path / "Ch06_remediated.pptx"
    pres_model.raw_presentation.save(out_file)
    assert out_file.exists()

    reloaded = reader.read(out_file)
    assert len(reloaded.slides) == initial_slide_count

    # Verify Slide 1 Shape 5 ("CHAPTER 6") text integrity
    s1 = reloaded.slides[0].raw_slide
    s1_shape5 = next((s for s in s1.shapes if s.shape_id == 5), None)
    assert s1_shape5 is not None
    assert "CHAPTER 6" in s1_shape5.text


# =============================================================================
# 2. Validation Deck: Sec1-5-TrigFunctions (2)original.pptx
# =============================================================================
def test_real_validation_trig_functions(tmp_path, reader, rule):
    deck_path = Path("tests/fixtures/Sec1-5-TrigFunctions (2)original.pptx")
    if not deck_path.exists():
        pytest.skip(f"{deck_path} not found.")

    pres_model = reader.read(deck_path)
    initial_slide_count = len(pres_model.slides)
    assert initial_slide_count > 0

    # Collect initial shape IDs and count
    initial_shape_counts = {
        s.slide_number: len(list(s.raw_slide.shapes))
        for s in pres_model.slides
    }

    audit_report = rule.generate_presentation_audit_report(pres_model)
    assert audit_report["total_slides_evaluated"] == initial_slide_count
    assert audit_report["visual_integrity_verified"] is True
    assert audit_report["failed_conversions"] == 0

    # Execute remediation
    for slide in pres_model.slides:
        issues = rule.evaluate(slide, pres_model)
        for issue in issues:
            if issue.behavior == RuleBehavior.SAFE_AUTO_REMEDIATE:
                action = rule.remediate_issue(slide, pres_model, issue)
                assert action.status in (RemediationStatus.APPLIED, RemediationStatus.SKIPPED)

    # Save and reload
    out_file = tmp_path / "Trig_remediated.pptx"
    pres_model.raw_presentation.save(out_file)
    reloaded = reader.read(out_file)
    assert len(reloaded.slides) == initial_slide_count

    # Shape count stability check (no lost shapes)
    for s in reloaded.slides:
        assert len(list(s.raw_slide.shapes)) >= initial_shape_counts[s.slide_number]


# =============================================================================
# 3. Validation Deck: Chapter 9.pptx (Cards, Badges '03', '04', Groups)
# =============================================================================
def test_real_validation_chapter_9(tmp_path, reader, rule):
    deck_path = Path("workspace_runs/job_20261008_082913_345db169/original/Chapter 9.pptx")
    if not deck_path.exists():
        pytest.skip(f"{deck_path} not found.")

    pres_model = reader.read(deck_path)
    initial_slide_count = len(pres_model.slides)
    assert initial_slide_count > 0

    audit_report = rule.generate_presentation_audit_report(pres_model)
    assert audit_report["total_slides_evaluated"] == initial_slide_count
    assert audit_report["visual_integrity_verified"] is True
    assert audit_report["failed_conversions"] == 0

    # Check that numeric card badges like '03', '04' are never promoted to titles
    for obj in audit_report["objects"]:
        preview = obj.get("text_preview", "").strip()
        if preview in ("01", "02", "03", "04", "05", "#01", "#02", "#03"):
            assert obj["action"] in (ConversionAction.PRESERVE_EXISTING, ConversionAction.SKIP_UNSAFE)
            assert obj["proposed_placeholder_type"] != "title"

    # Remediate presentation
    applied = 0
    for slide in pres_model.slides:
        issues = rule.evaluate(slide, pres_model)
        for issue in issues:
            if issue.behavior == RuleBehavior.SAFE_AUTO_REMEDIATE:
                action = rule.remediate_issue(slide, pres_model, issue)
                if action.status == RemediationStatus.APPLIED:
                    applied += 1

    # Save and reload
    out_file = tmp_path / "Chapter9_remediated.pptx"
    pres_model.raw_presentation.save(out_file)
    assert out_file.exists()

    reloaded = reader.read(out_file)
    assert len(reloaded.slides) == initial_slide_count
