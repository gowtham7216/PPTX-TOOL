"""
Tests for SlideGuard PresentationReader non-mutating XML parsing.

Verifies that reading/analyzing a presentation never mutates the underlying
OOXML DOM by injecting empty <a:pPr/>, <a:rPr/>, <a:solidFill/>, or <p:txBody/> elements.
"""

from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from pptx_engine.reader.presentation_reader import PresentationReader
from pptx_engine.ooxml.xml_namespaces import NAMESPACES
from models.enums import ListType


def test_reader_does_not_mutate_absent_ppr_and_absent_solidfill(tmp_path):
    """Verifies that paragraphs without pPr and runs without solidFill remain unchanged."""
    prs_path = tmp_path / "test_non_mutating.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Add textbox with text
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Hello non-mutating world"

    prs.save(str(prs_path))

    # Inspect slide XML before reading
    prs_before = Presentation(str(prs_path))
    slide_xml_before = etree.tostring(prs_before.slides[0]._element, encoding="unicode")

    # Read with PresentationReader
    reader = PresentationReader()
    model = reader.read(prs_path)

    # Underlying presentation attached to model
    prs_after = model.raw_presentation
    slide_xml_after = etree.tostring(prs_after.slides[0]._element, encoding="unicode")

    # Assert XML did not change at all during reading
    assert slide_xml_before == slide_xml_after

    # Verify model correctly read defaults
    s1 = model.slides[0]
    sh1 = s1.shapes[0]
    assert len(sh1.paragraphs) == 1
    p_model = sh1.paragraphs[0]
    assert p_model.level == 0
    assert p_model.list_type == ListType.NONE
    assert len(p_model.runs) >= 1
    r_model = p_model.runs[0]
    assert r_model.font_color_rgb is None  # Inherited theme color, no explicit RGB


def test_reader_preserves_explicit_rgb_and_theme_colors(tmp_path):
    """Verifies explicit RGB colors and theme-inherited colors are handled safely."""
    prs_path = tmp_path / "test_colors.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
    tf = txBox.text_frame

    # Paragraph 1: explicit red run
    p1 = tf.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "Red Run"
    r1.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    # Paragraph 2: run with no explicit color (theme-inherited)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Theme Run"

    # Paragraph 3: explicit level 2 with bold
    p3 = tf.add_paragraph()
    p3.level = 2
    r3 = p3.add_run()
    r3.text = "Level 2 Bold Run"
    r3.font.bold = True
    r3.font.size = Pt(20)

    prs.save(str(prs_path))

    # Snapshot XML before reading
    prs_before = Presentation(str(prs_path))
    slide_xml_before = etree.tostring(prs_before.slides[0]._element, encoding="unicode")

    reader = PresentationReader()
    model = reader.read(prs_path)

    prs_after = model.raw_presentation
    slide_xml_after = etree.tostring(prs_after.slides[0]._element, encoding="unicode")

    # Assert 100% XML immutability during reading
    assert slide_xml_before == slide_xml_after

    # Verify model extracted properties accurately
    s1 = model.slides[0]
    sh1 = s1.shapes[0]
    assert len(sh1.paragraphs) == 3

    # Check paragraph 1
    p1_m = sh1.paragraphs[0]
    assert p1_m.level == 0
    assert p1_m.runs[0].font_color_rgb == "FF0000"

    # Check paragraph 2
    p2_m = sh1.paragraphs[1]
    assert p2_m.level == 0
    assert p2_m.runs[0].font_color_rgb is None  # Preserved as None without adding solidFill

    # Check paragraph 3
    p3_m = sh1.paragraphs[2]
    assert p3_m.level == 2
    assert p3_m.runs[0].font_bold is True
    assert p3_m.runs[0].font_size_pt == 20.0


def test_reader_does_not_mutate_shape_without_txbody(tmp_path):
    """Verifies that an AutoShape without a txBody element is not mutated by reading."""
    prs_path = tmp_path / "test_no_txbody.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Add rectangle shape (which python-pptx creates with txBody, then remove txBody to test)
    rect = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
    tx_body = rect._sp.find(f"{{{NAMESPACES['p']}}}txBody")
    if tx_body is not None:
        rect._sp.remove(tx_body)

    prs.save(str(prs_path))

    prs_before = Presentation(str(prs_path))
    slide_xml_before = etree.tostring(prs_before.slides[0]._element, encoding="unicode")
    # Verify txBody was indeed absent
    assert prs_before.slides[0].shapes[0]._sp.find(f"{{{NAMESPACES['p']}}}txBody") is None

    reader = PresentationReader()
    model = reader.read(prs_path)

    prs_after = model.raw_presentation
    slide_xml_after = etree.tostring(prs_after.slides[0]._element, encoding="unicode")

    # Still absent! Not mutated!
    assert prs_after.slides[0].shapes[0]._sp.find(f"{{{NAMESPACES['p']}}}txBody") is None
    assert slide_xml_before == slide_xml_after

    # Model reports empty text
    sh = model.slides[0].shapes[0]
    assert sh.text_content == ""
    assert len(sh.paragraphs) == 0


def test_reader_immutability_on_real_deck_ch06():
    """Verifies that reading real deck Chapter 6 does not mutate any of its 50 slides."""
    deck_path = Path("tests/fixtures/Ch06_GLN_Instructor_AU.pptx")
    if not deck_path.exists():
        return

    prs_before = Presentation(str(deck_path))
    xmls_before = [etree.tostring(s._element, encoding="unicode") for s in prs_before.slides]

    reader = PresentationReader()
    model = reader.read(deck_path)

    prs_after = model.raw_presentation
    xmls_after = [etree.tostring(s._element, encoding="unicode") for s in prs_after.slides]

    assert len(xmls_before) == len(xmls_after) == 50
    for idx, (b, a) in enumerate(zip(xmls_before, xmls_after), start=1):
        assert b == a, f"Slide {idx} XML was mutated during reading!"


def test_reader_immutability_on_real_deck_trig():
    """Verifies that reading real deck Trig does not mutate any of its 19 slides."""
    deck_path = Path("tests/fixtures/Sec1-5-TrigFunctions (2)original.pptx")
    if not deck_path.exists():
        return

    prs_before = Presentation(str(deck_path))
    xmls_before = [etree.tostring(s._element, encoding="unicode") for s in prs_before.slides]

    reader = PresentationReader()
    model = reader.read(deck_path)

    prs_after = model.raw_presentation
    xmls_after = [etree.tostring(s._element, encoding="unicode") for s in prs_after.slides]

    assert len(xmls_before) == len(xmls_after) == 19
    for idx, (b, a) in enumerate(zip(xmls_before, xmls_after), start=1):
        assert b == a, f"Slide {idx} XML was mutated during reading!"
