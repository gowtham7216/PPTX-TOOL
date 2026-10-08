"""
SlideGuard Acceptance and Unit Tests for Task 2: Placeholder Conversion
powered by Gemini 3.5 Runtime Semantic Analysis.

Covers:
1. Runtime Model Configuration & Strict Separation:
   - SlideGuard runtime model is Gemini 3.5 (preferred) or Gemini 3.1 (fallback)
   - Gemini 3.8 Flash High is NOT selected as SlideGuard runtime model for Task 2
   - Antigravity Coding Model (Gemini 3.8 Flash High) remains untouched and decoupled
2. Candidate Detection: loose textboxes, native placeholders, picture/table/chart, empty
3. Context Extraction: zero PPTX filesystem paths or binaries transmitted
4. Gemini 3.5 Semantic Decision: strict JSON parsing and schema validation
5. Deterministic Safety Gate:
   - High confidence (>= 0.85) safe auto-conversion
   - Medium confidence (0.70-0.849) confirmation gate
   - Low confidence (< 0.70) preservation & manual review
   - Animation target preservation (never modified)
   - Single title constraint (no duplicate titles)
   - Specialized and empty placeholder preservation
6. Deterministic PPTX/XML Conversion: shape IDs, geometry, text preserved
7. Atomic Rollback: verified safety rollback if post-conversion validation fails
8. Model Traceability in remediation audit records
"""

import json
import os
from pathlib import Path
import pytest
from unittest.mock import MagicMock
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
from lxml import etree

from app.config import AppConfig
from engine.intelligence.gemini_assistant import (
    GeminiAssistant,
    resolve_runtime_model,
    ANTIGRAVITY_CODING_MODEL,
    DEFAULT_RUNTIME_MODEL,
    FALLBACK_RUNTIME_MODEL,
)
from rules.placeholders.classifier import (
    PlaceholderClassifier,
    PlaceholderCategory,
    PlaceholderStatus,
    PlaceholderAnalysis,
    _safe_xpath,
)
from rules.placeholders.transformer import (
    PlaceholderTransformer,
    ConversionResult,
)
from rules.placeholders.rule import PlaceholderRule
from pptx_engine.reader.presentation_reader import PresentationReader
from models.slide import SlideModel
from models.presentation import PresentationModel
from models.issue import Issue
from models.enums import RemediationStatus, RuleBehavior


# =============================================================================
# 1. RUNTIME MODEL CONFIGURATION & SEPARATION ACCEPTANCE TESTS
# =============================================================================

def test_acceptance_runtime_model_configured_to_gemini_35():
    """
    Acceptance Test: Confirms that SlideGuard runtime Gemini client is configured
    to Gemini 3.5 (preferred), that Gemini 3.8 is NOT selected for Task 2,
    and that Antigravity coding model remains untouched.
    """
    # 1. Config level check
    config = AppConfig.load_from_env()
    assert config.gemini_model == "gemini-3.5", (
        f"Expected default SlideGuard runtime model 'gemini-3.5', got '{config.gemini_model}'"
    )

    # 2. Assistant level check
    assistant = GeminiAssistant()
    assert assistant.model_name == "gemini-3.5", (
        f"Expected GeminiAssistant default runtime model 'gemini-3.5', got '{assistant.model_name}'"
    )

    # 3. Rule / Classifier level check for Task 2 Placeholder Conversion
    rule = PlaceholderRule()
    assert rule.classifier.assistant.model_name == "gemini-3.5", (
        f"Expected PlaceholderRule runtime model 'gemini-3.5', got '{rule.classifier.assistant.model_name}'"
    )

    # 4. Confirm Gemini 3.8 is NOT selected as SlideGuard runtime model
    assert rule.classifier.assistant.model_name != "gemini-3.8-flash-high"
    assert rule.classifier.assistant.model_name != "gemini-3.8"

    # 5. Confirm resolve_runtime_model prefers Gemini 3.5 and disallows coding model
    resolved_default = resolve_runtime_model()
    assert resolved_default == "gemini-3.5"

    resolved_filtered = resolve_runtime_model(preferred_model="gemini-3.8-flash-high")
    assert resolved_filtered == "gemini-3.5", "Coding model gemini-3.8-flash-high must be filtered out for runtime"

    resolved_fallback = resolve_runtime_model(preferred_model="gemini-3.1")
    assert resolved_fallback == "gemini-3.1"

    # 6. Confirm Antigravity Coding Model remains untouched
    assert ANTIGRAVITY_CODING_MODEL == "Gemini 3.8 Flash High"
    # Ensure runtime model and coding model are completely independent
    assert assistant.model_name != ANTIGRAVITY_CODING_MODEL


# =============================================================================
# 2. CANDIDATE DETECTION & CONTEXT EXTRACTION TESTS
# =============================================================================

@pytest.fixture
def candidate_deck(tmp_path) -> Path:
    """Creates a presentation with diverse candidate shapes for Task 2."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Standalone Title candidate + Body textbox + Footer
    s1 = prs.slides.add_slide(blank_layout)
    # Loose title at top
    t1 = s1.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    t1.text_frame.text = "Standalone Slide Title Heading"
    # Loose body paragraph
    b1 = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(3))
    b1.text_frame.text = "This is a loose body text box describing key concepts."
    # Loose footer at bottom
    f1 = s1.shapes.add_textbox(Inches(1), Inches(7.0), Inches(4), Inches(0.4))
    f1.text_frame.text = "Confidential Company Document"

    # Slide 2: Native Title Placeholder + Native Body Placeholder + Loose Textbox
    # Layout 0 has Title (idx 0) and Subtitle (idx 1)
    s2 = prs.slides.add_slide(prs.slide_layouts[0])
    s2.placeholders[0].text = "Native Presentation Title"
    s2.placeholders[1].text = "Native Subtitle Text"
    extra_box = s2.shapes.add_textbox(Inches(1), Inches(0.6), Inches(5), Inches(0.8))
    extra_box.text_frame.text = "Loose textbox that should NOT become title"

    out_path = tmp_path / "candidate_deck.pptx"
    prs.save(str(out_path))
    return out_path


def test_candidate_detection_and_context_privacy(candidate_deck):
    """Verifies candidate detection gathers shapes and context extraction contains no paths."""
    reader = PresentationReader()
    pres = reader.read(candidate_deck)

    classifier = PlaceholderClassifier()
    s1 = pres.slides[0]

    # Intercept prompt to check privacy
    recorded_prompts = []
    def intercepting_client(prompt: str) -> str:
        recorded_prompts.append(prompt)
        return json.dumps({
            "slide_number": 1,
            "confidence": 0.95,
            "rationale": "Test rationale",
            "decisions": [
                {
                    "shape_id": 2,
                    "candidate_type": "STANDALONE_TITLE",
                    "target_placeholder_type": "title",
                    "action": "CONVERT_TO_TITLE",
                    "confidence": 0.95,
                    "reason": "Top heading"
                },
                {
                    "shape_id": 3,
                    "candidate_type": "BODY_TEXTBOX",
                    "target_placeholder_type": "none",
                    "action": "PRESERVE",
                    "confidence": 0.90,
                    "reason": "Body content"
                },
                {
                    "shape_id": 4,
                    "candidate_type": "FOOTER",
                    "target_placeholder_type": "ftr",
                    "action": "PRESERVE",
                    "confidence": 0.90,
                    "reason": "Footer note"
                }
            ]
        })

    classifier.assistant = GeminiAssistant(client=intercepting_client, model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(
        raw_slide=s1.raw_slide,
        slide_number=1,
        slide_height_pt=pres.slide_height_pt,
        slide_width_pt=pres.slide_width_pt
    )

    assert len(recorded_prompts) == 1
    prompt_text = recorded_prompts[0]

    # Context Privacy Check: NEVER expose PPTX filesystem paths
    assert str(candidate_deck) not in prompt_text
    assert "candidate_deck.pptx" not in prompt_text
    assert "C:\\" not in prompt_text and "c:/" not in prompt_text

    # Candidate Detection Check: prompt must include candidate coordinates and text
    assert "Standalone Slide Title Heading" in prompt_text
    assert "geometry" in prompt_text


# =============================================================================
# 3. GEMINI 3.5 SEMANTIC DECISION WITH MOCK CLIENT
# =============================================================================

def test_gemini_35_decision_strict_json():
    """Validates that GeminiAssistant parses strict JSON decisions using runtime model gemini-3.5."""
    mock_response = json.dumps({
        "slide_number": 3,
        "confidence": 0.92,
        "rationale": "Slide contains one unlinked title box and one body placeholder.",
        "decisions": [
            {
                "shape_id": 10,
                "candidate_type": "STANDALONE_TITLE",
                "target_placeholder_type": "title",
                "action": "CONVERT_TO_TITLE",
                "confidence": 0.94,
                "reason": "Primary heading at top of slide lacks native placeholder semantics."
            },
            {
                "shape_id": 11,
                "candidate_type": "BODY_PLACEHOLDER",
                "target_placeholder_type": "none",
                "action": "MATERIALIZE_BODY",
                "confidence": 0.90,
                "reason": "Body placeholder with accessible content."
            }
        ]
    })

    assistant = GeminiAssistant(client=lambda p: mock_response, model_name="gemini-3.5")
    context = {
        "slide_number": 3,
        "slide_dimensions": {"width_pt": 720.0, "height_pt": 540.0},
        "existing_native_title_id": None,
        "candidates": [
            {"shape_id": 10, "name": "TitleBox", "current_ph_type": None, "has_text": True},
            {"shape_id": 11, "name": "ContentPlaceholder", "current_ph_type": "body", "has_text": True}
        ]
    }

    result = assistant.analyze_placeholder_conversion(context)
    assert result["model_used"] == "gemini-3.5"
    assert result["confidence"] == 0.92
    assert len(result["decisions"]) == 2
    assert result["decisions"][0]["action"] == "CONVERT_TO_TITLE"
    assert result["decisions"][1]["action"] == "MATERIALIZE_BODY"


# =============================================================================
# 4. DETERMINISTIC SAFETY GATE TESTS
# =============================================================================

def test_safety_gate_single_title_rule(candidate_deck):
    """
    Safety Gate Rule: If slide already has a native title placeholder,
    loose text boxes must NEVER be converted to title (must be PRESERVED).
    """
    reader = PresentationReader()
    pres = reader.read(candidate_deck)
    classifier = PlaceholderClassifier()

    # Slide 2 already has a native title (shape_id 2)
    s2 = pres.slides[1]

    # Even if Gemini were to propose CONVERT_TO_TITLE on the loose box:
    mock_ai = {
        "slide_number": 2,
        "confidence": 0.90,
        "rationale": "Attempting duplicate title",
        "decisions": [
            {"shape_id": 2, "action": "PRESERVE", "confidence": 0.95, "reason": "Native Title"},
            {"shape_id": 3, "action": "PRESERVE", "confidence": 0.95, "reason": "Native Subtitle"},
            {"shape_id": 4, "action": "CONVERT_TO_TITLE", "confidence": 0.90, "reason": "Duplicate candidate"}
        ]
    }
    classifier.assistant = GeminiAssistant(client=lambda p: json.dumps(mock_ai), model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(
        raw_slide=s2.raw_slide,
        slide_number=2,
        slide_height_pt=pres.slide_height_pt,
        slide_width_pt=pres.slide_width_pt
    )

    # Confirm only native title and subtitle are placeholders, no duplicate title
    title_analyses = [a for a in analyses if a.category == PlaceholderCategory.STANDALONE_TITLE_BOX]
    assert len(title_analyses) == 0, "Safety Gate must disallow second title when native title exists"


def test_safety_gate_animated_placeholder_preserved(tmp_path):
    """
    Safety Gate Rule: Placeholders targeted by animation sequences must NEVER
    be auto-converted (forced MANUAL_REVIEW / PRESERVE).
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "Animated Title"
    shape_id = box.shape_id

    # Add animation target spTgt
    sp_elem = box._element
    slide_elem = s._element
    timing_xml = parse_xml(
        f'<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        f'  <p:tnLst><p:par><p:cTn id="1"><p:childTnLst><p:seq>'
        f'    <p:cTn id="2"><p:childTnLst><p:par>'
        f'      <p:cTn id="3">'
        f'        <p:spTgt spid="{shape_id}"/>'
        f'      </p:cTn>'
        f'    </p:par></p:childTnLst></p:cTn>'
        f'  </p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'</p:timing>'
    )
    slide_elem.append(timing_xml)

    classifier = PlaceholderClassifier()
    # Mock AI tries to convert to title
    mock_ai = {
        "slide_number": 1,
        "confidence": 0.95,
        "rationale": "High confidence title",
        "decisions": [
            {"shape_id": shape_id, "action": "CONVERT_TO_TITLE", "confidence": 0.95, "reason": "Visual title"}
        ]
    }
    classifier.assistant = GeminiAssistant(client=lambda p: json.dumps(mock_ai), model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(
        raw_slide=s,
        slide_number=1,
        animation_target_ids=[shape_id]
    )

    # Animated shape must NOT be SAFE_TO_CONVERT
    animated_analyses = [a for a in analyses if a.shape_id == shape_id]
    if animated_analyses:
        assert animated_analyses[0].status == PlaceholderStatus.MANUAL_REVIEW
        assert not animated_analyses[0].is_safe_to_convert


def test_safety_gate_low_confidence_preserved(tmp_path):
    """
    Safety Gate Rule: Low confidence (< 0.70) decisions must force PRESERVE
    and flag for manual review.
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "Ambiguous Heading"
    shape_id = box.shape_id

    classifier = PlaceholderClassifier()
    # Mock AI with low confidence
    mock_ai = {
        "slide_number": 1,
        "confidence": 0.55,
        "rationale": "Uncertain structure",
        "decisions": [
            {"shape_id": shape_id, "action": "CONVERT_TO_TITLE", "confidence": 0.55, "reason": "Ambiguous"}
        ]
    }
    classifier.assistant = GeminiAssistant(client=lambda p: json.dumps(mock_ai), model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(
        raw_slide=s,
        slide_number=1
    )

    # Must NOT be SAFE_TO_CONVERT
    assert not any(a.is_safe_to_convert for a in analyses)


# =============================================================================
# 5. ATOMIC ROLLBACK ON INTEGRITY FAILURE
# =============================================================================

def test_atomic_rollback_on_unsafe_corruption(tmp_path):
    """
    Validates that PlaceholderTransformer executes atomic rollback
    if post-conversion verification detects an integrity failure.
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "Rollback Candidate Title"
    shape_id = box.shape_id

    transformer = PlaceholderTransformer()
    original_xml = etree.tostring(s._element)

    # Patch verify_integrity to simulate a failure
    transformer.verify_integrity = MagicMock(return_value=(False, ["Simulated corrupted shape ID loss"]))

    result = transformer.convert_slide(
        slide_elem=s._element,
        slide_number=1,
        raw_slide=s,
        approved_shape_ids=[shape_id]
    )

    # Must report failure
    assert not result.success
    assert any("Simulated corrupted" in err for err in result.errors)
    assert any("Atomic rollback executed" in err for err in result.errors)

    # Slide XML must be restored to original
    current_xml = etree.tostring(s._element)
    assert current_xml == original_xml, "Slide XML must match pre-conversion state after rollback"


# =============================================================================
# 6. MODEL TRACEABILITY IN REMEDIATION AUDIT
# =============================================================================

def test_model_traceability_in_remediation_action(tmp_path):
    """
    Verifies that PlaceholderRule.remediate_issue records model traceability:
    - Antigravity Coding Model: Gemini 3.8 Flash High
    - SlideGuard Runtime Model: Gemini 3.5
    """
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    box.text_frame.text = "Traceability Title Test"

    rule = PlaceholderRule()
    slide_model = SlideModel(slide_id=1, slide_number=1, raw_slide=s)
    pres_model = PresentationModel(file_path=tmp_path / "dummy.pptx")

    issues = rule.evaluate(slide_model, pres_model)
    assert len(issues) == 1
    issue = issues[0]

    action = rule.remediate_issue(slide_model, pres_model, issue)
    assert action.status == RemediationStatus.APPLIED

    # Check model traceability
    traceability = action.remediated_state.get("model_traceability", {})
    assert traceability.get("antigravity_coding_model") == "Gemini 3.8 Flash High"
    assert traceability.get("slideguard_runtime_model") == "Gemini 3.5"
    assert traceability.get("runtime_model_identifier") == "gemini-3.5"


# =============================================================================
# 7. TASK 2B: 15 SPECIFIC UNIT & ACCEPTANCE TESTS
# =============================================================================

def test_1_card_badge_03_near_top_preserved():
    """1. '03' card badge near top of slide -> PRESERVE."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # Place '03' badge at y = 1.0 inch (top region)
    b3 = s.shapes.add_textbox(Inches(5), Inches(1.0), Inches(0.5), Inches(0.5))
    b3.text_frame.text = "03"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)
    badge_analysis = next((a for a in analyses if a.shape_id == b3.shape_id), None)
    # The badge must NOT be STANDALONE_TITLE_BOX and must NOT be SAFE_TO_CONVERT
    if badge_analysis:
        assert badge_analysis.category != PlaceholderCategory.STANDALONE_TITLE_BOX
        assert badge_analysis.status == PlaceholderStatus.PRESERVED


def test_2_card_badge_03_before_real_title_in_xml():
    """2. '03' card badge appears before real title in XML -> real title wins."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # Add '03' first (lower shape ID / first in XML)
    b3 = s.shapes.add_textbox(Inches(5), Inches(1.5), Inches(0.5), Inches(0.5))
    b3.text_frame.text = "03"
    # Add real title second (higher shape ID / second in XML)
    t = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    t.text_frame.text = "Case Study: The Invisible Stop Sign Attack"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    title_analysis = next((a for a in analyses if a.shape_id == t.shape_id), None)
    badge_analysis = next((a for a in analyses if a.shape_id == b3.shape_id), None)

    assert title_analysis is not None
    assert title_analysis.category == PlaceholderCategory.STANDALONE_TITLE_BOX
    assert title_analysis.status == PlaceholderStatus.SAFE_TO_CONVERT

    if badge_analysis:
        assert badge_analysis.category != PlaceholderCategory.STANDALONE_TITLE_BOX
        assert badge_analysis.status == PlaceholderStatus.PRESERVED


def test_3_three_repeated_cards_with_badges_none_become_title():
    """3. Three repeated cards with 01/02/03 -> none become TITLE."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    b1 = s.shapes.add_textbox(Inches(1), Inches(1.5), Inches(0.5), Inches(0.5))
    b1.text_frame.text = "01"
    b2 = s.shapes.add_textbox(Inches(3), Inches(1.5), Inches(0.5), Inches(0.5))
    b2.text_frame.text = "02"
    b3 = s.shapes.add_textbox(Inches(5), Inches(1.5), Inches(0.5), Inches(0.5))
    b3.text_frame.text = "03"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    badge_ids = {b1.shape_id, b2.shape_id, b3.shape_id}
    for a in analyses:
        if a.shape_id in badge_ids:
            assert a.category != PlaceholderCategory.STANDALONE_TITLE_BOX
            assert a.status == PlaceholderStatus.PRESERVED


def test_4_real_title_after_card_badge_in_xml_selected():
    """4. Real title after card badge in XML -> real title selected."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    badge = s.shapes.add_textbox(Inches(2), Inches(1.0), Inches(0.4), Inches(0.4))
    badge.text_frame.text = "01"
    heading = s.shapes.add_textbox(Inches(1), Inches(0.4), Inches(8), Inches(1.0))
    heading.text_frame.text = "Introduction to Adversarial Attacks on AI"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    heading_analysis = next((a for a in analyses if a.shape_id == heading.shape_id), None)
    assert heading_analysis is not None
    assert heading_analysis.category == PlaceholderCategory.STANDALONE_TITLE_BOX
    assert heading_analysis.status == PlaceholderStatus.SAFE_TO_CONVERT


def test_5_real_title_below_28_percent_height_detected():
    """5. Real title below 28% slide height -> still detected."""
    prs = Presentation()
    prs.slide_height = Inches(7.5) # 540 pt
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # Placed at y = 4.5 inches (60% down the slide, like Slide 1)
    low_title = s.shapes.add_textbox(Inches(1), Inches(4.5), Inches(8), Inches(1.2))
    low_title.text_frame.text = "Chapter 9 - The Cat-and-Mouse Game: Adversarial Attacks on AI"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    title_analysis = next((a for a in analyses if a.shape_id == low_title.shape_id), None)
    assert title_analysis is not None
    assert title_analysis.category == PlaceholderCategory.STANDALONE_TITLE_BOX
    assert title_analysis.status == PlaceholderStatus.SAFE_TO_CONVERT


def test_6_slide_with_no_reliable_title_no_forced_conversion():
    """6. Slide with no reliable title -> no forced conversion."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # Only tiny fragments and numbers
    box1 = s.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
    box1.text_frame.text = "..."
    box2 = s.shapes.add_textbox(Inches(2), Inches(2), Inches(1), Inches(1))
    box2.text_frame.text = "#"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    assert not any(a.category == PlaceholderCategory.STANDALONE_TITLE_BOX for a in analyses)
    assert not any(a.is_safe_to_convert for a in analyses)


def test_7_existing_native_title_preserve():
    """7. Existing native title -> preserve, no extra title created."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.placeholders[0].text = "Native Presentation Heading"
    extra_box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    extra_box.text_frame.text = "Another Prominent Textbox"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    # Native title must be preserved
    native_analysis = next((a for a in analyses if a.category == PlaceholderCategory.TITLE or a.placeholder_type in ("title", "ctrTitle")), None)
    assert native_analysis is not None
    assert native_analysis.status == PlaceholderStatus.PRESERVED

    # Extra textbox must NOT be converted to title
    extra_analysis = next((a for a in analyses if a.shape_id == extra_box.shape_id), None)
    if extra_analysis:
        assert extra_analysis.category != PlaceholderCategory.STANDALONE_TITLE_BOX


def test_8_multiple_plausible_titles_gemini_ranking():
    """8. Multiple plausible titles -> Gemini ranks candidates and selects winner."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box_sub = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(4), Inches(0.8))
    box_sub.text_frame.text = "Subheading Topic A"
    box_main = s.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(1.2))
    box_main.text_frame.text = "Comprehensive Security Architecture for AI Models"

    mock_gemini = {
        "slide_number": 1,
        "selected_title_shape_id": box_main.shape_id,
        "overall_confidence": 0.96,
        "reason": "Shape box_main is the comprehensive primary slide title.",
        "candidates": [
            {"shape_id": box_sub.shape_id, "role": "SUBTITLE", "action": "PRESERVE", "confidence": 0.90},
            {"shape_id": box_main.shape_id, "role": "TITLE", "action": "CONVERT_TO_TITLE", "confidence": 0.96}
        ]
    }
    classifier = PlaceholderClassifier()
    classifier.assistant = GeminiAssistant(client=lambda p: json.dumps(mock_gemini), model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    main_analysis = next((a for a in analyses if a.shape_id == box_main.shape_id), None)
    assert main_analysis is not None
    assert main_analysis.category == PlaceholderCategory.STANDALONE_TITLE_BOX
    assert main_analysis.status == PlaceholderStatus.SAFE_TO_CONVERT


def test_9_gemini_incorrectly_selects_badge_safety_gate_rejects():
    """9. Gemini incorrectly selects badge -> deterministic safety gate rejects it."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    badge = s.shapes.add_textbox(Inches(5), Inches(1.0), Inches(0.5), Inches(0.5))
    badge.text_frame.text = "03"
    real_title = s.shapes.add_textbox(Inches(1), Inches(0.4), Inches(8), Inches(1))
    real_title.text_frame.text = "Tesla Autopilot Phantom Lane Attack"

    # Gemini mistakenly selects badge shape_id
    bad_gemini = {
        "slide_number": 1,
        "selected_title_shape_id": badge.shape_id,
        "overall_confidence": 0.90,
        "reason": "Mistaken badge selection",
        "candidates": [
            {"shape_id": badge.shape_id, "role": "TITLE", "action": "CONVERT_TO_TITLE", "confidence": 0.90},
            {"shape_id": real_title.shape_id, "role": "BODY_TEXTBOX", "action": "PRESERVE", "confidence": 0.90}
        ]
    }
    classifier = PlaceholderClassifier()
    classifier.assistant = GeminiAssistant(client=lambda p: json.dumps(bad_gemini), model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    # Safety gate MUST have rejected badge
    badge_analysis = next((a for a in analyses if a.shape_id == badge.shape_id), None)
    if badge_analysis:
        assert badge_analysis.category != PlaceholderCategory.STANDALONE_TITLE_BOX
        assert badge_analysis.status == PlaceholderStatus.PRESERVED

    # Real title should be selected instead
    title_analysis = next((a for a in analyses if a.shape_id == real_title.shape_id), None)
    assert title_analysis is not None
    assert title_analysis.category == PlaceholderCategory.STANDALONE_TITLE_BOX
    assert title_analysis.status == PlaceholderStatus.SAFE_TO_CONVERT


def test_10_gemini_unavailable_preserve_and_safe_fallback():
    """10. Gemini unavailable -> preserve or safe deterministic ranking."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    t.text_frame.text = "Reliable Slide Title When AI Offline"

    def failing_client(prompt: str) -> str:
        raise ConnectionError("Gemini service unreachable")

    classifier = PlaceholderClassifier()
    classifier.assistant = GeminiAssistant(client=failing_client, model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    # Deterministic fallback handles it safely
    title_analysis = next((a for a in analyses if a.shape_id == t.shape_id), None)
    assert title_analysis is not None
    assert title_analysis.status == PlaceholderStatus.SAFE_TO_CONVERT


def test_11_gemini_invalid_json_fallback_safe():
    """11. Gemini invalid JSON -> fallback gracefully without corruption."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    t.text_frame.text = "Title Handled Despite Malformed JSON"

    classifier = PlaceholderClassifier()
    classifier.assistant = GeminiAssistant(client=lambda p: "THIS IS NOT VALID JSON", model_name="gemini-3.5")
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    title_analysis = next((a for a in analyses if a.shape_id == t.shape_id), None)
    assert title_analysis is not None
    assert title_analysis.status == PlaceholderStatus.SAFE_TO_CONVERT


def test_12_animated_candidate_preserve_manual_review():
    """12. Animated candidate -> preserve/manual review, never converted to title."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    t = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    t.text_frame.text = "Animated Title Candidate"
    anim_id = t.shape_id

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(
        raw_slide=s,
        slide_number=1,
        animation_target_ids=[anim_id]
    )

    t_analysis = next((a for a in analyses if a.shape_id == anim_id), None)
    assert t_analysis is not None
    assert t_analysis.is_animated
    assert t_analysis.status == PlaceholderStatus.MANUAL_REVIEW
    assert not t_analysis.is_safe_to_convert


def test_13_repeated_design_labels_preserved():
    """13. Repeated design labels -> preserve."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    l1 = s.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.4))
    l1.text_frame.text = "Step 1:"
    l2 = s.shapes.add_textbox(Inches(3), Inches(1), Inches(1), Inches(0.4))
    l2.text_frame.text = "Step 2:"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    for a in analyses:
        if a.shape_id in (l1.shape_id, l2.shape_id):
            assert a.category != PlaceholderCategory.STANDALONE_TITLE_BOX
            assert a.status == PlaceholderStatus.PRESERVED


def test_14_numeric_only_text_never_title():
    """14. Numeric-only text -> never TITLE."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    num_box = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(1), Inches(0.5))
    num_box.text_frame.text = "42"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    num_analysis = next((a for a in analyses if a.shape_id == num_box.shape_id), None)
    if num_analysis:
        assert num_analysis.category != PlaceholderCategory.STANDALONE_TITLE_BOX
        assert num_analysis.status == PlaceholderStatus.PRESERVED


def test_15_short_decorative_text_never_title():
    """15. Short decorative text -> never TITLE."""
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    arrow = s.shapes.add_textbox(Inches(1), Inches(0.5), Inches(0.5), Inches(0.5))
    arrow.text_frame.text = "-->"

    classifier = PlaceholderClassifier()
    analyses = classifier.analyze_slide_placeholders(raw_slide=s, slide_number=1)

    arrow_analysis = next((a for a in analyses if a.shape_id == arrow.shape_id), None)
    if arrow_analysis:
        assert arrow_analysis.category != PlaceholderCategory.STANDALONE_TITLE_BOX
        assert arrow_analysis.status == PlaceholderStatus.PRESERVED
