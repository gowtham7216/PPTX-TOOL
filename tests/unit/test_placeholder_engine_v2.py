from lxml import etree
import pytest

from rules.placeholders.placeholder_engine_v2 import (
    PlaceholderEngineV2,
    mark_decorative,
    normalize_content_placeholder_names,
)

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"p": P, "a": A}


def make_shape(shape_id=2, ph=None, name=None):
    sp = etree.Element(f"{{{P}}}sp", nsmap={"p": P, "a": A})
    nv = etree.SubElement(sp, f"{{{P}}}nvSpPr")
    shape_name = name or f"Shape {shape_id}"
    etree.SubElement(nv, f"{{{P}}}cNvPr", id=str(shape_id), name=shape_name)
    etree.SubElement(nv, f"{{{P}}}cNvSpPr")
    nvpr = etree.SubElement(nv, f"{{{P}}}nvPr")
    if ph is not None:
        etree.SubElement(nvpr, f"{{{P}}}ph", attrib=ph)
    sppr = etree.SubElement(sp, f"{{{P}}}spPr")
    xfrm = etree.SubElement(sppr, f"{{{A}}}xfrm")
    etree.SubElement(xfrm, f"{{{A}}}off", x="10", y="20")
    etree.SubElement(xfrm, f"{{{A}}}ext", cx="30", cy="40")
    tx = etree.SubElement(sp, f"{{{P}}}txBody")
    body_pr = etree.SubElement(tx, f"{{{A}}}bodyPr")
    etree.SubElement(body_pr, f"{{{A}}}spAutoFit")
    etree.SubElement(tx, f"{{{A}}}lstStyle")
    para = etree.SubElement(tx, f"{{{A}}}p")
    run = etree.SubElement(para, f"{{{A}}}r")
    etree.SubElement(run, f"{{{A}}}t").text = "Keep this text"
    return sp


def make_tree(*shapes):
    tree = etree.Element(f"{{{P}}}spTree", nsmap={"p": P, "a": A})
    for shape in shapes:
        tree.append(shape)
    return tree


def test_adds_native_placeholder_and_preserves_text_geometry_and_id():
    engine = PlaceholderEngineV2()
    shape = make_shape(7)
    before_text = "".join(shape.xpath(".//a:t/text()", namespaces=NS))
    before_xfrm = etree.tostring(shape.xpath("./p:spPr/a:xfrm", namespaces=NS)[0])
    result = engine.ensure_placeholder(shape, "title", make_tree(shape))
    ph = shape.xpath("./p:nvSpPr/p:nvPr/p:ph", namespaces=NS)[0]
    assert result.inserted is True
    assert ph.get("type") == "title"
    assert ph.get("idx") is None
    assert shape.xpath("./p:nvSpPr/p:cNvPr/@id", namespaces=NS) == ["7"]
    assert "".join(shape.xpath(".//a:t/text()", namespaces=NS)) == before_text
    assert etree.tostring(shape.xpath("./p:spPr/a:xfrm", namespaces=NS)[0]) == before_xfrm


def test_existing_placeholder_is_preserved_without_rewriting_type_or_index():
    engine = PlaceholderEngineV2()
    shape = make_shape(4, {"type": "body", "idx": "8"})
    before = etree.tostring(shape)
    result = engine.ensure_placeholder(shape, "title", make_tree(shape))
    assert result.inserted is False
    assert result.placeholder_type == "body"
    assert result.placeholder_idx == "8"
    assert etree.tostring(shape) == before


def test_body_index_avoids_collision_with_existing_placeholder_indices():
    engine = PlaceholderEngineV2()
    existing = make_shape(1, {"type": "body", "idx": "1"})
    target = make_shape(2)
    tree = make_tree(existing, target)
    result = engine.ensure_placeholder(target, "body", tree)
    assert result.inserted is True
    assert result.placeholder_idx == "2"


def test_subtitle_uses_next_free_index_and_not_hardcoded_one():
    engine = PlaceholderEngineV2()
    existing = make_shape(1, {"type": "body", "idx": "1"})
    target = make_shape(2)
    tree = make_tree(existing, target)
    result = engine.ensure_placeholder(target, "subTitle", tree)
    assert result.placeholder_idx == "2"


def test_rejects_unsupported_placeholder_type():
    with pytest.raises(ValueError, match="Unsupported placeholder type"):
        PlaceholderEngineV2().ensure_placeholder(make_shape(), "madeUp")


def test_freeze_paragraph_properties_and_autofit_on_body_conversion():
    engine = PlaceholderEngineV2()
    shape = make_shape(5)
    result = engine.ensure_placeholder(shape, "body")
    assert result.inserted is True

    # 1. <a:buNone/> is injected in paragraph properties to isolate from master layout bullets
    bu_none = shape.xpath(".//a:p/a:pPr/a:buNone", namespaces=NS)
    assert len(bu_none) == 1

    # 2. Line spacing and paragraph spacing are set
    ln_spc = shape.xpath(".//a:p/a:pPr/a:lnSpc/a:spcPct/@val", namespaces=NS)
    assert ln_spc == ["100000"]
    spc_bef = shape.xpath(".//a:p/a:pPr/a:spcBef/a:spcPts/@val", namespaces=NS)
    assert spc_bef == ["0"]
    spc_aft = shape.xpath(".//a:p/a:pPr/a:spcAft/a:spcPts/@val", namespaces=NS)
    assert spc_aft == ["0"]

    # 3. Autofit is frozen to <a:noAutofit/>
    no_autofit = shape.xpath(".//a:bodyPr/a:noAutofit", namespaces=NS)
    assert len(no_autofit) == 1
    sp_autofit = shape.xpath(".//a:bodyPr/a:spAutoFit", namespaces=NS)
    assert len(sp_autofit) == 0

    # 4. Run size is backfilled
    r_sz = shape.xpath(".//a:p/a:r/a:rPr/@sz", namespaces=NS)
    assert len(r_sz) == 1 and r_sz[0] == "1800"


def test_preserves_real_bullet_lists():
    engine = PlaceholderEngineV2()
    shape = make_shape(6)
    para = shape.xpath(".//a:p", namespaces=NS)[0]
    ppr = etree.SubElement(para, f"{{{A}}}pPr")
    etree.SubElement(ppr, f"{{{A}}}buChar", char="•")

    result = engine.ensure_placeholder(shape, "body")
    assert result.inserted is True

    # buChar must remain and buNone must NOT be inserted
    bu_char = shape.xpath(".//a:p/a:pPr/a:buChar", namespaces=NS)
    assert len(bu_char) == 1
    bu_none = shape.xpath(".//a:p/a:pPr/a:buNone", namespaces=NS)
    assert len(bu_none) == 0


def test_content_placeholder_renaming_and_title_name_preservation():
    engine = PlaceholderEngineV2()

    # Body placeholder is renamed to "Content Placeholder {idx}"
    body_shape = make_shape(10, name="TextBox 10")
    engine.ensure_placeholder(body_shape, "body", preferred_idx="3")
    name = body_shape.xpath(".//p:cNvPr/@name", namespaces=NS)[0]
    assert name == "Content Placeholder 3"

    # Title placeholder preserves original name
    title_shape = make_shape(11, name="Custom Title Heading")
    engine.ensure_placeholder(title_shape, "title")
    title_name = title_shape.xpath(".//p:cNvPr/@name", namespaces=NS)[0]
    assert title_name == "Custom Title Heading"


def test_decorative_marking_preserved():
    engine = PlaceholderEngineV2()
    shape = make_shape(12)
    mark_decorative(shape)

    result = engine.ensure_placeholder(shape, "body")
    assert result.inserted is False
    assert result.reason == "decorative-shape-preserved"
    # No p:ph added
    assert len(shape.xpath(".//p:ph", namespaces=NS)) == 0


def test_normalize_content_placeholder_names():
    sh1 = make_shape(1, ph={"type": "body", "idx": "4"}, name="Old TextBox 1")
    sh2 = make_shape(2, ph={"type": "title"}, name="Slide Title")
    sh3 = make_shape(3, ph={"type": "body", "idx": "5"}, name="Content Placeholder 5")
    tree = make_tree(sh1, sh2, sh3)

    renamed = normalize_content_placeholder_names(tree)
    assert renamed == [("1", "Content Placeholder 4")]
    assert sh1.xpath(".//p:cNvPr/@name", namespaces=NS)[0] == "Content Placeholder 4"
    assert sh2.xpath(".//p:cNvPr/@name", namespaces=NS)[0] == "Slide Title"
    assert sh3.xpath(".//p:cNvPr/@name", namespaces=NS)[0] == "Content Placeholder 5"


def make_empty_shape(shape_id=99, name=None):
    sp = etree.Element(f"{{{P}}}sp", nsmap={"p": P, "a": A})
    nv = etree.SubElement(sp, f"{{{P}}}nvSpPr")
    shape_name = name or f"Shape {shape_id}"
    etree.SubElement(nv, f"{{{P}}}cNvPr", id=str(shape_id), name=shape_name)
    etree.SubElement(nv, f"{{{P}}}cNvSpPr")
    etree.SubElement(nv, f"{{{P}}}nvPr")
    sppr = etree.SubElement(sp, f"{{{P}}}spPr")
    xfrm = etree.SubElement(sppr, f"{{{A}}}xfrm")
    etree.SubElement(xfrm, f"{{{A}}}off", x="10", y="20")
    etree.SubElement(xfrm, f"{{{A}}}ext", cx="30", cy="40")
    return sp


def test_skip_empty_text_shapes_leaves_empty_shapes_unconverted():
    from pptx import Presentation
    from rules.placeholders.transformer import PlaceholderTransformer
    transformer = PlaceholderTransformer()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    empty_shape = make_empty_shape(88, "Empty Background Tile")
    text_shape = make_shape(89, name="Real Content Card")
    slide.shapes._spTree.append(empty_shape)
    slide.shapes._spTree.append(text_shape)

    # When skip_empty_text=True: empty shape preserved untouched without <p:ph>, content shape converted
    res = transformer.convert_slide(slide._element, slide_number=1, raw_slide=slide, skip_empty_text=True)
    assert res.success is True
    assert 88 not in res.converted_shape_ids
    assert len(empty_shape.xpath(".//p:ph", namespaces=NS)) == 0
    assert 89 in res.converted_shape_ids
    assert len(text_shape.xpath(".//p:ph", namespaces=NS)) == 1


def test_font_preservation_across_shapes_and_remediation():
    """Validates that fonts standardize to Times New Roman by default, and fall back to Calibri upon content overflow."""
    from pptx import Presentation
    from rules.placeholders.transformer import PlaceholderTransformer
    transformer = PlaceholderTransformer()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Create shape 101 with ample box (470pt x 118pt) -> Times New Roman fits
    tnr_shape = make_shape(101, name="TNR Box")
    ext1 = tnr_shape.xpath(".//a:ext", namespaces=NS)[0]
    ext1.set("cx", "6000000")
    ext1.set("cy", "1500000")
    r1 = tnr_shape.xpath(".//a:r", namespaces=NS)[0]
    rpr1 = etree.Element(f"{{{A}}}rPr")
    lat1 = etree.SubElement(rpr1, f"{{{A}}}latin")
    lat1.set("typeface", "Arial")  # Original font is Arial
    r1.insert(0, rpr1)
    
    # Create shape 102 with tight box (100pt x 25pt) and multi-line text -> Times New Roman overflows -> falls back to Calibri
    cal_shape = make_shape(102, name="Calibri Box")
    ext2 = cal_shape.xpath(".//a:ext", namespaces=NS)[0]
    ext2.set("cx", "1270000")
    ext2.set("cy", "320000")
    tx2 = cal_shape.xpath(".//a:t", namespaces=NS)[0]
    tx2.text = "This is a longer paragraph that requires several lines of text to render inside a tight bounding box."
    r2 = cal_shape.xpath(".//a:r", namespaces=NS)[0]
    rpr2 = etree.Element(f"{{{A}}}rPr")
    lat2 = etree.SubElement(rpr2, f"{{{A}}}latin")
    lat2.set("typeface", "Figtree")  # Original font is Figtree
    r2.insert(0, rpr2)
    
    slide.shapes._spTree.append(tnr_shape)
    slide.shapes._spTree.append(cal_shape)
    res = transformer.convert_slide(slide._element, slide_number=1, raw_slide=slide)
    assert res.success is True
    assert 101 in res.converted_shape_ids
    assert 102 in res.converted_shape_ids
    
    # Verify Times New Roman selected for shape 101 (fits box)
    tnr_fonts = tnr_shape.xpath(".//a:latin/@typeface", namespaces=NS)
    assert "Times New Roman" in tnr_fonts
    
    # Verify Calibri selected for shape 102 (overflow fallback)
    cal_fonts = cal_shape.xpath(".//a:latin/@typeface", namespaces=NS)
    assert "Calibri" in cal_fonts


def test_font_preservation_mode():
    """Validates that when standardize_fonts=False, the original font is strictly preserved."""
    from rules.placeholders.placeholder_engine_v2 import preserve_and_backfill_run_fonts
    
    shape = make_shape(103, name="Preserve Shape")
    r = shape.xpath(".//a:r", namespaces=NS)[0]
    rpr = etree.Element(f"{{{A}}}rPr")
    lat = etree.SubElement(rpr, f"{{{A}}}latin")
    lat.set("typeface", "Courier New")
    r.insert(0, rpr)
    
    applied = preserve_and_backfill_run_fonts(shape, standardize_fonts=False)
    assert applied == "Courier New"
    fonts = shape.xpath(".//a:latin/@typeface", namespaces=NS)
    assert "Courier New" in fonts

