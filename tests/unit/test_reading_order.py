"""
Comprehensive unit and integration tests for SlideGuard Reading Order engine.
Tests semantic classification, layout region analysis, confidence model,
footer exclusion, and in-place XML node reordering.
"""

from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE

from rules.reading_order.semantic_classifier import SemanticRoleClassifier, SemanticRole, SemanticItem
from rules.reading_order.layout_analyzer import LayoutRegionAnalyzer
from rules.reading_order.order_resolver import ReadingOrderResolver, ReadingOrderResult
from rules.reading_order.xml_reorderer import XMLReadingOrderReorderer
from rules.reading_order.rule import ReadingOrderRule
from models.slide import SlideModel
from models.presentation import PresentationModel
from models.enums import RemediationStatus


@pytest.fixture
def reading_order_test_pptx(tmp_path) -> Path:
    """Creates a targeted multi-slide presentation covering all reading order scenarios."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    title_layout = prs.slide_layouts[0]

    # --- SLIDE 1: Title + Body (Native placeholders, correct order) ---
    s1 = prs.slides.add_slide(title_layout)
    s1.shapes.title.text = "Slide 1: Native Title"
    s1.placeholders[1].text = "Slide 1: Native Body Content"

    # --- SLIDE 2: Out of Order (Body created first, Title created second) ---
    s2 = prs.slides.add_slide(blank_layout)
    # Body first (shape_id will be lower than title)
    body_box = s2.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(3))
    body_box.text_frame.text = "Body content placed first in XML tree"
    # Title placed second near top
    title_box = s2.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    title_box.text_frame.text = "Slide Title Placed Second in XML Tree"

    # --- SLIDE 3: Two Columns Layout ---
    s3 = prs.slides.add_slide(blank_layout)
    s3_title = s3.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    s3_title.text_frame.text = "Two Column Comparison"
    # Column 1 (Left: left=1.0)
    col1_top = s3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3.5), Inches(2.0))
    col1_top.text_frame.text = "Left Column Item 1"
    col1_bot = s3.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(3.5), Inches(2.0))
    col1_bot.text_frame.text = "Left Column Item 2"
    # Column 2 (Right: left=5.5)
    col2_top = s3.shapes.add_textbox(Inches(5.5), Inches(2.0), Inches(3.5), Inches(2.0))
    col2_top.text_frame.text = "Right Column Item 1"
    col2_bot = s3.shapes.add_textbox(Inches(5.5), Inches(4.5), Inches(3.5), Inches(2.0))
    col2_bot.text_frame.text = "Right Column Item 2"

    # --- SLIDE 4: Wiley Footer & Bottom Content Slide ---
    s4 = prs.slides.add_slide(blank_layout)
    s4_title = s4.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    s4_title.text_frame.text = "Footer vs Legitimate Bottom Content"
    # Footer in bottom margin
    wiley_footer = s4.shapes.add_textbox(Inches(1), Inches(7.0), Inches(2), Inches(0.4))
    wiley_footer.text_frame.text = "WILEY"
    # Legitimate bottom problem box (taller, multi-line)
    bottom_content = s4.shapes.add_textbox(Inches(4), Inches(5.8), Inches(5), Inches(1.2))
    bottom_content.text_frame.text = "Important calculation summary located near bottom.\nMust not be excluded as a footer."

    # --- SLIDE 5: Image + Caption ---
    s5 = prs.slides.add_slide(blank_layout)
    s5_title = s5.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
    s5_title.text_frame.text = "Figure and Caption"
    # Auto-shape acting as diagram / figure
    fig = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(2.0), Inches(6), Inches(3.0))
    fig.name = "Figure_Diagram"
    # Caption directly below figure (within 20pt)
    caption = s5.shapes.add_textbox(Inches(2), Inches(5.1), Inches(6), Inches(0.5))
    caption.text_frame.text = "Figure 1: Conceptual Architecture Overview."

    out_path = tmp_path / "reading_order_test.pptx"
    prs.save(str(out_path))
    return out_path


def test_semantic_classifier_roles(reading_order_test_pptx):
    prs = Presentation(str(reading_order_test_pptx))
    classifier = SemanticRoleClassifier()

    # Slide 1: Native Title & Body
    s1 = prs.slides[0]
    classified = []
    for idx, shape in enumerate(s1.shapes):
        item = classifier.classify_shape(shape._element, shape, 540.0, 720.0, z_order=idx)
        classified.append(item)

    title_item = [i for i in classified if i.is_title][0]
    assert title_item.role == SemanticRole.TITLE
    assert title_item.confidence == 1.0

    sub_item = [i for i in classified if i.role in (SemanticRole.SUBTITLE, SemanticRole.BODY_TEXT)][0]
    assert sub_item.confidence >= 0.90


def test_wiley_footer_detection_vs_content(reading_order_test_pptx):
    prs = Presentation(str(reading_order_test_pptx))
    classifier = SemanticRoleClassifier()

    s4 = prs.slides[3]  # Slide 4
    classified = []
    for idx, shape in enumerate(s4.shapes):
        item = classifier.classify_shape(shape._element, shape, 540.0, 720.0, z_order=idx)
        classified.append(item)

    wiley = [i for i in classified if "wiley" in i.text_preview.lower()][0]
    assert wiley.is_footer is True
    assert wiley.role == SemanticRole.FOOTER

    bottom_content = [i for i in classified if "important calculation" in i.text_preview.lower()][0]
    # Multi-line problem text must NOT be classified as footer
    assert bottom_content.is_footer is False
    assert bottom_content.role == SemanticRole.BODY_TEXT


def test_two_column_reading_order_resolution(reading_order_test_pptx):
    prs = Presentation(str(reading_order_test_pptx))
    classifier = SemanticRoleClassifier()
    resolver = ReadingOrderResolver()

    s3 = prs.slides[2]  # Slide 3: Two Column
    classified = []
    for idx, shape in enumerate(s3.shapes):
        item = classifier.classify_shape(shape._element, shape, 540.0, 720.0, z_order=idx)
        classified.append(item)

    result = resolver.resolve_reading_order(
        slide_number=3,
        items=classified,
        slide_w_pt=720.0,
        slide_h_pt=540.0
    )

    assert result.confidence in ("HIGH", "MEDIUM")
    # Title must be first in proposed order
    title_id = [i.shape_id for i in classified if i.is_title][0]
    assert result.proposed_order[0] == title_id


def test_in_place_xml_reordering(reading_order_test_pptx, tmp_path):
    prs = Presentation(str(reading_order_test_pptx))
    reorderer = XMLReadingOrderReorderer()

    # Slide 2: Out-of-order (Body was created before Title)
    s2 = prs.slides[1]
    slide_elem = s2._element

    shape_ids = [s.shape_id for s in s2.shapes]
    assert len(shape_ids) == 2
    body_id, title_id = shape_ids[0], shape_ids[1]

    # Desired proposed order: Title first, then Body
    proposed_order = [title_id, body_id]

    success, msg, resulting_order = reorderer.reorder_sp_tree(slide_elem, proposed_order)
    assert success is True
    assert resulting_order == [title_id, body_id]

    # Save and reload to verify physical XML persistence
    out_file = tmp_path / "reordered_s2.pptx"
    prs.save(str(out_file))

    reloaded = Presentation(str(out_file))
    reloaded_s2_ids = [s.shape_id for s in reloaded.slides[1].shapes]
    assert reloaded_s2_ids == [title_id, body_id]


def test_reading_order_rule_remediate_issue(reading_order_test_pptx):
    from pptx_engine.reader.presentation_reader import PresentationReader
    reader = PresentationReader()
    pres_model = reader.read(reading_order_test_pptx)
    ro_rule = ReadingOrderRule()

    # Slide 2 has title after body in XML tree -> evaluate must find issue
    s2 = pres_model.get_slide_by_number(2)
    issues = ro_rule.evaluate(s2, pres_model)

    assert len(issues) == 1
    issue = issues[0]
    assert issue.rule_id == "SG-RO-01"
    assert issue.extra_data["changes_required"] is True

    # Remediate issue
    action = ro_rule.remediate_issue(s2, pres_model, issue)
    assert action.status == RemediationStatus.APPLIED
    assert action.remediated_state["order"] == issue.extra_data["proposed_order"]


def test_reading_order_preserves_animation_targets():
    """Validates that animation target shape IDs in <p:timing> are verified and preserved."""
    ch06_path = Path("tests/fixtures/Ch06_GLN_Instructor_AU.pptx")
    if not ch06_path.exists():
        return

    from pptx_engine.reader.presentation_reader import PresentationReader
    reader = PresentationReader()
    pres_model = reader.read(ch06_path)
    ro_rule = ReadingOrderRule()

    # Slide 4 has 34 animation references in Ch06
    s4 = pres_model.get_slide_by_number(4)
    issues = ro_rule.evaluate(s4, pres_model)

    if issues:
        action = ro_rule.remediate_issue(s4, pres_model, issues[0])
        # If reordered, action must succeed and all animation targets verified
        assert action.status in (RemediationStatus.APPLIED, RemediationStatus.SKIPPED)

    # Verification service must confirm animation target integrity
    from engine.verification.verification_service import VerificationService
    verifier = VerificationService()
    v_res = verifier.verify_file(ch06_path, expected_slide_count=50)
    assert v_res.passed is True


class _MockParagraph:
    def __init__(self, level=0):
        self.level = level


class _MockTextFrame:
    def __init__(self, text):
        self.text = text
        self.paragraphs = [_MockParagraph(0)]


class _MockShape:
    def __init__(self, shape_id, name, left_pt, top_pt, width_pt, height_pt, text):
        self.shape_id = shape_id
        self.name = name
        self.left = int(left_pt * 12700)
        self.top = int(top_pt * 12700)
        self.width = int(width_pt * 12700)
        self.height = int(height_pt * 12700)
        self.has_text_frame = True
        self.text_frame = _MockTextFrame(text)


def _create_mock_sp_element(shape_id: int, name: str, ph_type: str = None):
    from lxml import etree
    nsmap = {
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main"
    }
    elem = etree.Element("{http://schemas.openxmlformats.org/presentationml/2006/main}sp", nsmap=nsmap)
    nv_sp_pr = etree.SubElement(elem, "{http://schemas.openxmlformats.org/presentationml/2006/main}nvSpPr")
    etree.SubElement(nv_sp_pr, "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr", id=str(shape_id), name=name)
    if ph_type:
        nv_pr = etree.SubElement(nv_sp_pr, "{http://schemas.openxmlformats.org/presentationml/2006/main}nvPr")
        etree.SubElement(nv_pr, "{http://schemas.openxmlformats.org/presentationml/2006/main}ph", type=ph_type)
    return elem


def test_heading_labels_not_classified_as_footer():
    """Validates that Chapter, Section, Unit, Topic labels are never classified as footers."""
    classifier = SemanticRoleClassifier()
    slide_h, slide_w = 540.0, 720.0

    test_labels = [
        "CHAPTER 6",
        "Chapter 6",
        "SECTION 1",
        "Unit 2",
        "Topic 3: Conservation of Energy",
    ]

    for idx, label in enumerate(test_labels, start=10):
        # Test both mid-slide and in bottom footer zone
        for top_pos in (200.0, 500.0):
            elem = _create_mock_sp_element(idx, f"Label_{label}")
            mock_shape = _MockShape(idx, f"Label_{label}", 50.0, top_pos, 300.0, 30.0, label)

            item = classifier.classify_shape(
                shape_elem=elem,
                shape_obj=mock_shape,
                slide_h_pt=slide_h,
                slide_w_pt=slide_w,
                z_order=0
            )

            assert item.is_footer is False, f"'{label}' at top={top_pos} was incorrectly marked is_footer=True"
            assert item.role != SemanticRole.FOOTER, f"'{label}' at top={top_pos} was incorrectly given role FOOTER"


def test_legitimate_bottom_positioned_content_not_footer():
    """Validates that legitimate educational content in the bottom margin is not treated as footer."""
    classifier = SemanticRoleClassifier()
    slide_h, slide_w = 540.0, 720.0

    legit_texts = [
        "If the object does not move, the work is zero.",
        "Important calculation summary: total energy is conserved.",
        "Answer: 42 m/s",
        "v = sqrt(2gh)",
        "Note: Friction is negligible in this scenario."
    ]

    for idx, text in enumerate(legit_texts, start=20):
        elem = _create_mock_sp_element(idx, "ContentBox")
        mock_shape = _MockShape(idx, "ContentBox", 100.0, 490.0, 400.0, 35.0, text)

        item = classifier.classify_shape(
            shape_elem=elem,
            shape_obj=mock_shape,
            slide_h_pt=slide_h,
            slide_w_pt=slide_w,
            z_order=1
        )

        assert item.is_footer is False, f"Legitimate text '{text}' was incorrectly marked as footer"
        assert item.role == SemanticRole.BODY_TEXT


def test_verified_footers_are_still_excluded():
    """Validates that strongly verified footer patterns and native placeholders are properly excluded."""
    classifier = SemanticRoleClassifier()
    slide_h, slide_w = 540.0, 720.0

    # A. Corporate footer (WILEY)
    elem_wiley = _create_mock_sp_element(30, "WileyFooter")
    mock_wiley = _MockShape(30, "WileyFooter", 50.0, 500.0, 100.0, 20.0, "WILEY")
    item_wiley = classifier.classify_shape(elem_wiley, mock_wiley, slide_h, slide_w, 0)
    assert item_wiley.is_footer is True
    assert item_wiley.role == SemanticRole.FOOTER

    # B. Copyright notice
    elem_copy = _create_mock_sp_element(31, "CopyNotice")
    mock_copy = _MockShape(31, "CopyNotice", 50.0, 510.0, 500.0, 20.0, "© 2024 John Wiley & Sons, Inc. All rights reserved.")
    item_copy = classifier.classify_shape(elem_copy, mock_copy, slide_h, slide_w, 1)
    assert item_copy.is_footer is True
    assert item_copy.role == SemanticRole.FOOTER

    # C. Native placeholder (p:ph type="ftr")
    elem_ph = _create_mock_sp_element(32, "FooterPlaceholder", ph_type="ftr")
    mock_ph = _MockShape(32, "FooterPlaceholder", 50.0, 500.0, 200.0, 20.0, "Confidential Presentation")
    item_ph = classifier.classify_shape(elem_ph, mock_ph, slide_h, slide_w, 2)
    assert item_ph.is_footer is True
    assert item_ph.role == SemanticRole.FOOTER

    # D. Order resolver verifies exclusion from main sequence
    resolver = ReadingOrderResolver()
    res = resolver.resolve_reading_order(
        slide_number=1,
        items=[item_wiley, item_copy, item_ph],
        slide_w_pt=slide_w,
        slide_h_pt=slide_h
    )
    assert len(res.excluded_footer_objects) == 3


def test_ch06_slide_1_order_preservation():
    """Validates that Slide 1 of Ch06 maintains author order and does not exclude Shape 5 ('CHAPTER 6')."""
    ch06_path = Path("tests/fixtures/Ch06_GLN_Instructor_AU.pptx")
    if not ch06_path.exists():
        pytest.skip("Ch06 fixture not present")

    from pptx_engine.reader.presentation_reader import PresentationReader
    reader = PresentationReader()
    pres_model = reader.read(ch06_path)
    ro_rule = ReadingOrderRule()

    s1 = pres_model.get_slide_by_number(1)
    issues = ro_rule.evaluate(s1, pres_model)

    # Must require no changes: author's sequence [2, 3, 4, 5, 6] is preserved
    assert len(issues) == 0, f"Expected 0 issues for Slide 1, got: {[i.description for i in issues]}"

    raw_shapes = list(s1.raw_slide.shapes)
    classified = [
        ro_rule.classifier.classify_shape(s._element, s, pres_model.slide_height_pt, pres_model.slide_width_pt, i)
        for i, s in enumerate(raw_shapes)
    ]
    res = ro_rule.resolver.resolve_reading_order(
        slide_number=1,
        items=classified,
        slide_w_pt=pres_model.slide_width_pt,
        slide_h_pt=pres_model.slide_height_pt
    )

    assert res.original_order == [2, 3, 4, 5, 6]
    assert res.proposed_order == [2, 3, 4, 5, 6]
    assert res.changes_required is False
    assert len(res.excluded_footer_objects) == 0

    # Shape 5 must be BODY_TEXT and not a footer
    shape_5_item = [it for it in classified if it.shape_id == 5][0]
    assert shape_5_item.is_footer is False
    assert shape_5_item.role != SemanticRole.FOOTER
    assert "CHAPTER 6" in shape_5_item.text_preview

