"""
SlideGuard QA Tests: Relationship & Hyperlink Integrity (Phase 10).
Verifies that:
- External hyperlinks remain valid and point to correct URLs.
- Image relationships remain connected to media parts without dangling references.
- Animation target shape IDs remain intact without broken timing linkages.
"""

import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches

from engine.qa.qa_service import QAService
from pptx_engine.reader.presentation_reader import PresentationReader
from pptx_engine.writer.presentation_writer import PresentationWriter
from tests.qa.synthetic_profiles import create_profile_animation, _get_temp_png


def test_hyperlinks_preserved_across_pipeline():
    """Verify external hyperlinks on shapes are intact after write and reload."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        source_pptx = tmp / "hyperlinks.pptx"

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        p = tb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = "Click to visit W3C"
        run.hyperlink.address = "https://www.w3.org/WAI/"
        prs.save(str(source_pptx))

        # Read and re-write
        reader = PresentationReader()
        writer = PresentationWriter()
        pres = reader.read(source_pptx)

        out_pptx = tmp / "hyperlinks_out.pptx"
        writer.write(pres, out_pptx)

        # Inspect relationships in output package
        qa_service = QAService()
        rel_check = qa_service.verify_relationships_and_hyperlinks(source_pptx, out_pptx)
        assert rel_check["preserved"] is True
        assert rel_check["candidate_hyperlinks"] >= 1


def test_image_media_relationships_intact():
    """Verify image relationships remain valid and uncorrupted in package."""
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        img_pptx = tmp / "images.pptx"
        png_path = _get_temp_png(tmp)

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(png_path), Inches(1), Inches(1), Inches(3), Inches(2))
        prs.save(str(img_pptx))

        reader = PresentationReader()
        writer = PresentationWriter()
        pres = reader.read(img_pptx)

        out_pptx = tmp / "images_out.pptx"
        writer.write(pres, out_pptx)

        qa_service = QAService()
        rel_check = qa_service.verify_relationships_and_hyperlinks(img_pptx, out_pptx)
        assert rel_check["preserved"] is True
        assert rel_check["candidate_media_rels"] >= 1


def test_animation_timing_and_target_ids_preserved_untouched():
    """Verify animation <p:timing> and targeted shape IDs are completely preserved without remediation."""
    qa_service = QAService()
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        anim_pptx = tmp / "animations.pptx"
        create_profile_animation(anim_pptx)

        res = qa_service.run_comprehensive_qa(anim_pptx, profile_name="ANIMATION_PROFILE")
        assert res.animation_preservation_status.value == "PASS"
        assert res.integrity_status.value == "PASS"
