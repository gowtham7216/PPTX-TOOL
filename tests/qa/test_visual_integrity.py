"""
SlideGuard QA Tests: Visual & Design Preservation (Phase 10).
Verifies that:
- Slide dimensions, shape geometry (left, top, width, height), rotation, fills, borders,
  and typography are strictly preserved.
- Distinguishes intentional accessibility remediations from unintended visual drift.
- Confirms Design Preservation = PASS.
"""

import tempfile
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from engine.qa.qa_service import QAService
from engine.comparison.comparison_service import ComparisonService
from pptx_engine.reader.presentation_reader import PresentationReader
from pptx_engine.writer.presentation_writer import PresentationWriter
from tests.qa.synthetic_profiles import create_profile_design_heavy


def test_design_preservation_on_styled_presentation():
    """Verify geometry, styling, and colors remain identical when saved."""
    qa_service = QAService()
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        design_pptx = tmp / "design_styled.pptx"
        create_profile_design_heavy(design_pptx)

        res = qa_service.run_comprehensive_qa(design_pptx, profile_name="DESIGN_HEAVY")
        assert res.design_preservation_status.value == "PASS"
        assert res.integrity_status.value == "PASS"


def test_distinguishes_intentional_accessibility_changes_from_drift():
    """Verify intentional accessibility changes are not flagged as unintended drift."""
    reader = PresentationReader()
    comparer = ComparisonService()

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        test_pptx = tmp / "intentional_test.pptx"

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        tb.text_frame.text = "Original draft text"
        shape_id = tb.shape_id
        prs.save(str(test_pptx))

        orig = reader.read(test_pptx)
        rem = reader.read(test_pptx)

        # Simulate user approved AI text edit on shape
        rem_shape = rem.slides[0].get_shape_by_id(shape_id)
        rem_shape.text_content = "User curated accessible text"

        # If shape_id is declared intentional:
        result_approved = comparer.compare(orig, rem, intentional_shape_ids=[shape_id])
        assert result_approved.integrity_preserved is True
        assert result_approved.unintended_changes_count == 0

        # If shape_id is NOT approved (unintended drift):
        result_unapproved = comparer.compare(orig, rem, intentional_shape_ids=[])
        assert result_unapproved.integrity_preserved is False
        assert result_unapproved.unintended_changes_count == 1
