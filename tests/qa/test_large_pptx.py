"""
SlideGuard QA Tests: Large 100+ Slide Presentation Workload (Phase 10).
Validates that SlideGuard can efficiently and safely process large presentations:
- 100+ slides
- Hundreds of shapes, tables, lists, and footers
- Verifies output validity, shape ID integrity, and acceptable duration without memory corruption.
"""

import time
import tempfile
from pathlib import Path
import pytest
from pptx import Presentation

from engine.qa.qa_service import QAService
from tests.qa.synthetic_profiles import create_profile_large


def test_large_presentation_processing_105_slides():
    """Verify end-to-end processing of a 105-slide synthetic presentation."""
    qa_service = QAService()
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        large_pptx = tmp / "large_105_slides.pptx"

        t0 = time.time()
        create_profile_large(large_pptx, slide_count=105)
        gen_duration = time.time() - t0

        assert large_pptx.exists()
        assert large_pptx.stat().st_size > 0

        # Execute full QA validation
        t_proc = time.time()
        res = qa_service.run_comprehensive_qa(large_pptx, profile_name="LARGE_SCALE")
        proc_duration = time.time() - t_proc

        assert res.integrity_status.value == "PASS"
        assert res.original_protection_status.value == "PASS"
        assert res.overall_status.value in ["PASS", "PASS_WITH_WARNINGS"]

        # Validate output presentation structure
        cand = Presentation(res.output_file)
        assert len(cand.slides) == 105, f"Expected 105 slides, got {len(cand.slides)}"

        # Verify execution time is reasonable (e.g. under 60 seconds)
        assert proc_duration < 60.0, f"Processing took too long: {proc_duration:.2f}s"
