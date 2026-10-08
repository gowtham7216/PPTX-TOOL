"""
SlideGuard Presentation Verification Service.
Validates structural integrity, XML validity, non-corruption, and animation target preservation
of remediated PPTX files.
"""

from pathlib import Path
from typing import Optional, Set
import logging
from pptx import Presentation

from models.verification import VerificationResult

logger = logging.getLogger(__name__)


class VerificationService:
    """
    Validates that remediated PPTX files remain well-formed, readable,
    and structurally intact.
    """

    def verify_file(
        self,
        target_path: Path,
        expected_slide_count: int = 0,
        expected_animation_target_ids: Optional[Set[int]] = None
    ) -> VerificationResult:
        result = VerificationResult()
        path = Path(target_path).resolve()

        # Check 1: File existence and non-zero size
        if not path.exists():
            result.add_check("File Existence", False, f"File not found at {path}")
            return result

        file_size = path.stat().st_size
        if file_size == 0:
            result.add_check("File Size", False, "File is 0 bytes (corrupt write)")
            return result
        result.add_check("File Size", True, f"File size valid: {file_size} bytes")

        # Check 2: Load presentation via python-pptx / zip container
        try:
            prs = Presentation(str(path))
            result.add_check("OOXML Container Validity", True, "Successfully opened presentation XML package")
        except Exception as e:
            result.add_check("OOXML Container Validity", False, f"Presentation package is corrupt: {e}")
            logger.error("Verification failed for %s: %s", path, e)
            return result

        # Check 3: Slide count preservation
        actual_slide_count = len(prs.slides)
        if expected_slide_count > 0 and actual_slide_count != expected_slide_count:
            result.add_check(
                "Slide Count Preservation",
                False,
                f"Slide count mismatch: expected {expected_slide_count}, got {actual_slide_count}"
            )
        else:
            result.add_check(
                "Slide Count Preservation",
                True,
                f"Slide count verified: {actual_slide_count} slides"
            )

        # Check 4: Slide layout and shape integrity
        corrupt_slides = 0
        all_shape_ids = set()
        for s_idx, s in enumerate(prs.slides, start=1):
            try:
                _ = len(s.shapes)
                for sh in s.shapes:
                    all_shape_ids.add(sh.shape_id)
            except Exception as e:
                corrupt_slides += 1
                result.add_check(
                    f"Slide {s_idx} Shape Tree",
                    False,
                    f"Shape tree error: {e}",
                    slide_number=s_idx
                )

        if corrupt_slides == 0:
            result.add_check("Slide Shapes Well-Formed", True, "All slide shape trees loaded successfully")

        # Check 5: Animation target ID stability
        if expected_animation_target_ids:
            missing_anim_ids = expected_animation_target_ids - all_shape_ids
            if missing_anim_ids:
                result.add_check(
                    "Animation Target ID Preservation",
                    False,
                    f"Animation targets lost post-remediation: {missing_anim_ids}"
                )
            else:
                result.add_check(
                    "Animation Target ID Preservation",
                    True,
                    f"All {len(expected_animation_target_ids)} animated shape IDs verified intact."
                )

        return result
