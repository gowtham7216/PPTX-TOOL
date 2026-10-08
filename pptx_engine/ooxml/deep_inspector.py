"""
SlideGuard Deep OOXML Inspection Engine.
Provides read-only inspection, structural diagnostics, and element mapping
for PowerPoint PPTX presentations without modifying presentation files.
"""

from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
import json
import logging
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from pptx_engine.ooxml.xml_namespaces import NAMESPACES
from pptx_engine.ooxml.xml_utils import find_elements_by_xpath, get_shape_alt_text, get_shape_decorative_status
from pptx_engine.utilities.geometry import (
    emu_to_points,
    is_in_footer_region,
    calculate_group_child_slide_coordinates,
)

logger = logging.getLogger(__name__)


class OOXMLDeepInspector:
    """
    Read-only inspection engine for PowerPoint OOXML structures.
    Maps shape trees, group transformations, list metadata, placeholders,
    animation targets, and relationship graphs without mutating files.
    """

    def inspect_presentation(self, pptx_path: Path) -> Dict[str, Any]:
        """Performs comprehensive deep inspection of a presentation file."""
        target_path = Path(pptx_path).resolve()
        if not target_path.exists():
            raise FileNotFoundError(f"File not found: {target_path}")

        prs = Presentation(str(target_path))
        slide_w_emu = prs.slide_width
        slide_h_emu = prs.slide_height
        slide_w_pt = emu_to_points(slide_w_emu)
        slide_h_pt = emu_to_points(slide_h_emu)

        all_slides_data = []
        global_animation_target_ids = set()
        global_shape_ids = set()

        for s_idx, slide in enumerate(prs.slides, start=1):
            slide_data = self.inspect_slide(slide, s_idx, slide_w_pt, slide_h_pt)
            all_slides_data.append(slide_data)

            # Accumulate IDs
            for node in slide_data["flat_shape_list"]:
                global_shape_ids.add(node["shape_id"])
            for anim_id in slide_data["animation_info"]["target_shape_ids"]:
                global_animation_target_ids.add(anim_id)

        report = {
            "file_name": target_path.name,
            "total_slides": len(prs.slides),
            "slide_dimensions": {
                "width_emu": slide_w_emu,
                "height_emu": slide_h_emu,
                "width_pt": slide_w_pt,
                "height_pt": slide_h_pt,
            },
            "global_id_inventory": {
                "total_unique_shape_ids": len(global_shape_ids),
                "total_animated_shape_ids": len(global_animation_target_ids),
                "animated_shape_ids": sorted(list(global_animation_target_ids)),
            },
            "slides": all_slides_data,
        }
        return report

    def inspect_slide(self, slide, slide_number: int, slide_w_pt: float, slide_h_pt: float) -> Dict[str, Any]:
        """Extracts complete structural breakdown of an individual slide."""
        slide_elem = slide._element

        # 1. Slide Relationships
        relationships = self._inspect_slide_relationships(slide)

        # 2. Animations (<p:timing>)
        animation_info = self._inspect_animations(slide_elem, slide_number)

        # 3. Shape Tree & Nested Groups
        sp_tree = slide_elem.xpath(".//p:spTree")
        tree_nodes = []
        flat_list = []
        group_transforms = []

        if sp_tree:
            # Enumerate children directly under <p:spTree>
            child_idx = 0
            for child in sp_tree[0]:
                tag = child.tag.split("}")[-1]
                if tag in ("sp", "grpSp", "pic", "graphicFrame", "cxnSp"):
                    child_idx += 1
                    node = self._parse_tree_element(
                        child,
                        slide_number=slide_number,
                        tree_index=child_idx,
                        nesting_depth=0,
                        parent_group_id=None,
                        group_stack=[],
                        slide_h_pt=slide_h_pt,
                        animation_target_ids=animation_info["target_shape_ids"]
                    )
                    tree_nodes.append(node)
                    self._flatten_nodes(node, flat_list)
                    self._collect_group_transforms(node, group_transforms)

        # 4. Placeholders
        placeholders = self._inspect_placeholders(slide, slide_number)

        # 5. List Paragraphs
        lists = self._inspect_lists(slide, slide_number)

        # 6. Footers
        footers = [n for n in flat_list if n["footer_candidate_type"] != "NONE"]

        return {
            "slide_number": slide_number,
            "layout_name": getattr(slide.slide_layout, "name", "Unknown") if hasattr(slide, "slide_layout") else "Unknown",
            "shape_count": len(flat_list),
            "top_level_element_count": len(tree_nodes),
            "animation_info": animation_info,
            "relationships": relationships,
            "shape_tree": tree_nodes,
            "flat_shape_list": flat_list,
            "group_transforms": group_transforms,
            "placeholders": placeholders,
            "list_items": lists,
            "footer_candidates": footers,
        }

    def _parse_tree_element(
        self,
        elem,
        slide_number: int,
        tree_index: int,
        nesting_depth: int,
        parent_group_id: Optional[int],
        group_stack: List[Tuple[Tuple[float, float], Tuple[float, float], Tuple[float, float], Tuple[float, float]]],
        slide_h_pt: float,
        animation_target_ids: List[int]
    ) -> Dict[str, Any]:
        tag = elem.tag.split("}")[-1]

        # Extract Non-Visual Properties (p:cNvPr)
        cnv_prs = elem.xpath(".//p:cNvPr")
        shape_id = 0
        name = "Unknown"
        rel_ids = []
        if cnv_prs:
            shape_id = int(cnv_prs[0].get("id", "0"))
            name = cnv_prs[0].get("name", f"Shape_{shape_id}")

        # Check for embedded media / blip rels
        blips = elem.xpath(".//a:blip")
        for blip in blips:
            r_embed = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if r_embed:
                rel_ids.append(r_embed)

        # Hyperlinks
        hlinks = elem.xpath(".//a:hlinkClick")
        for hl in hlinks:
            r_id = hl.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            if r_id:
                rel_ids.append(r_id)

        # Geometric Transform (a:xfrm)
        off_pt, ext_pt, raw_transform_data = self._extract_element_geometry(elem, tag)

        # Compute real slide coordinates if inside group
        computed_slide_bbox = self._compute_slide_coordinates(off_pt, ext_pt, group_stack)

        # Text extraction
        text_content = ""
        tx_body = elem.xpath(".//p:txBody")
        if tx_body:
            paras = elem.xpath(".//a:p")
            text_lines = []
            for p in paras:
                p_text = "".join(r.text for r in p.xpath(".//a:t") if r.text)
                if p_text:
                    text_lines.append(p_text)
            text_content = "\n".join(text_lines)

        # Placeholder info (<p:ph>)
        is_ph = False
        ph_type = None
        ph_idx = None
        ph_elems = elem.xpath(".//p:ph")
        if ph_elems:
            is_ph = True
            ph_type = ph_elems[0].get("type", "body")
            ph_idx = ph_elems[0].get("idx")

        # Footer Candidate Classification
        footer_candidate_type = "NONE"
        if is_ph and ph_type == "ftr":
            footer_candidate_type = "ACTUAL_FOOTER_PLACEHOLDER"
        elif is_ph and ph_type == "sldNum":
            footer_candidate_type = "SLIDE_NUMBER_PLACEHOLDER"
        elif is_ph and ph_type == "dt":
            footer_candidate_type = "DATETIME_PLACEHOLDER"
        elif is_in_footer_region(computed_slide_bbox[1], computed_slide_bbox[3], slide_h_pt):
            if text_content and len(text_content.strip()) < 80:
                footer_candidate_type = "SPATIAL_BOTTOM_TEXT"

        # Determine object category
        obj_type = "AUTO_SHAPE"
        if tag == "grpSp":
            obj_type = "GROUP"
        elif tag == "pic":
            obj_type = "PICTURE"
        elif tag == "graphicFrame":
            obj_type = "TABLE" if elem.xpath(".//a:tbl") else "GRAPHIC_FRAME"
        elif is_ph:
            obj_type = "PLACEHOLDER"
        elif text_content and not is_ph:
            obj_type = "TEXT_BOX"

        node = {
            "slide_number": slide_number,
            "xml_tag": f"p:{tag}",
            "shape_id": shape_id,
            "name": name,
            "object_type": obj_type,
            "parent_group_id": parent_group_id,
            "nesting_depth": nesting_depth,
            "tree_index": tree_index,
            "local_bbox_pt": {
                "left": round(off_pt[0], 2),
                "top": round(off_pt[1], 2),
                "width": round(ext_pt[0], 2),
                "height": round(ext_pt[1], 2),
            },
            "slide_bbox_pt": {
                "left": round(computed_slide_bbox[0], 2),
                "top": round(computed_slide_bbox[1], 2),
                "width": round(computed_slide_bbox[2], 2),
                "height": round(computed_slide_bbox[3], 2),
            },
            "has_text": bool(text_content.strip()),
            "text_content": text_content[:120] if text_content else "",
            "is_placeholder": is_ph,
            "placeholder_type": ph_type,
            "placeholder_idx": ph_idx,
            "is_picture": tag == "pic",
            "is_table": obj_type == "TABLE",
            "is_group": tag == "grpSp",
            "footer_candidate_type": footer_candidate_type,
            "is_animated": shape_id in animation_target_ids,
            "relationship_ids": rel_ids,
            "raw_transform_data": raw_transform_data,
            "children": [],
        }

        # Recursive parsing of children if this is a group
        if tag == "grpSp":
            # Add this group's transform to group_stack
            new_group_stack = list(group_stack)
            if raw_transform_data:
                grp_frame = (
                    (raw_transform_data["off_x_emu"], raw_transform_data["off_y_emu"]),
                    (raw_transform_data["ext_cx_emu"], raw_transform_data["ext_cy_emu"]),
                    (raw_transform_data["chOff_x_emu"], raw_transform_data["chOff_y_emu"]),
                    (raw_transform_data["chExt_cx_emu"], raw_transform_data["chExt_cy_emu"]),
                )
                new_group_stack.append(grp_frame)

            child_sub_idx = 0
            for child_elem in elem:
                c_tag = child_elem.tag.split("}")[-1]
                if c_tag in ("sp", "grpSp", "pic", "graphicFrame", "cxnSp"):
                    child_sub_idx += 1
                    child_node = self._parse_tree_element(
                        child_elem,
                        slide_number=slide_number,
                        tree_index=child_sub_idx,
                        nesting_depth=nesting_depth + 1,
                        parent_group_id=shape_id,
                        group_stack=new_group_stack,
                        slide_h_pt=slide_h_pt,
                        animation_target_ids=animation_target_ids
                    )
                    node["children"].append(child_node)

        return node

    def _extract_element_geometry(self, elem, tag: str) -> Tuple[Tuple[float, float], Tuple[float, float], Optional[Dict[str, int]]]:
        """Extracts off, ext, chOff, and chExt from element's xfrm container."""
        if tag == "grpSp":
            xfrm = elem.xpath(".//p:grpSpPr/a:xfrm")
        else:
            xfrm = elem.xpath(".//p:spPr/a:xfrm") or elem.xpath(".//a:xfrm")

        if not xfrm:
            return (0.0, 0.0), (0.0, 0.0), None

        x_node = xfrm[0]
        off = x_node.xpath(".//a:off")
        ext = x_node.xpath(".//a:ext")
        ch_off = x_node.xpath(".//a:chOff")
        ch_ext = x_node.xpath(".//a:chExt")

        off_x = int(off[0].get("x", "0")) if off else 0
        off_y = int(off[0].get("y", "0")) if off else 0
        ext_cx = int(ext[0].get("cx", "0")) if ext else 0
        ext_cy = int(ext[0].get("cy", "0")) if ext else 0

        raw_data = {
            "off_x_emu": off_x,
            "off_y_emu": off_y,
            "ext_cx_emu": ext_cx,
            "ext_cy_emu": ext_cy,
            "chOff_x_emu": int(ch_off[0].get("x", "0")) if ch_off else off_x,
            "chOff_y_emu": int(ch_off[0].get("y", "0")) if ch_off else off_y,
            "chExt_cx_emu": int(ch_ext[0].get("cx", "0")) if ch_ext else ext_cx,
            "chExt_cy_emu": int(ch_ext[0].get("cy", "0")) if ch_ext else ext_cy,
        }

        return (emu_to_points(off_x), emu_to_points(off_y)), (emu_to_points(ext_cx), emu_to_points(ext_cy)), raw_data

    def _compute_slide_coordinates(
        self,
        local_off_pt: Tuple[float, float],
        local_ext_pt: Tuple[float, float],
        group_stack: List[Any]
    ) -> Tuple[float, float, float, float]:
        """Maps coordinates through nested group transforms into outer slide points."""
        if not group_stack:
            return (local_off_pt[0], local_off_pt[1], local_ext_pt[0], local_ext_pt[1])

        curr_x_emu = local_off_pt[0] * 12700
        curr_y_emu = local_off_pt[1] * 12700
        curr_w_emu = local_ext_pt[0] * 12700
        curr_h_emu = local_ext_pt[1] * 12700

        # Apply stack from innermost group to outer-most slide
        for grp_off, grp_ext, grp_ch_off, grp_ch_ext in reversed(group_stack):
            curr_x_emu, curr_y_emu, curr_w_emu, curr_h_emu = calculate_group_child_slide_coordinates(
                group_off=grp_off,
                group_ext=grp_ext,
                group_ch_off=grp_ch_off,
                group_ch_ext=grp_ch_ext,
                child_off=(curr_x_emu, curr_y_emu),
                child_ext=(curr_w_emu, curr_h_emu)
            )

        return (
            emu_to_points(curr_x_emu),
            emu_to_points(curr_y_emu),
            emu_to_points(curr_w_emu),
            emu_to_points(curr_h_emu)
        )

    def _inspect_animations(self, slide_elem, slide_number: int) -> Dict[str, Any]:
        """Inspects <p:timing> animations and targets."""
        timings = slide_elem.xpath(".//p:timing")
        has_timing = len(timings) > 0
        target_ids = []

        if has_timing:
            sp_tgts = timings[0].xpath(".//p:spTgt")
            for tgt in sp_tgts:
                spid = tgt.get("spid")
                if spid and spid.isdigit():
                    target_ids.append(int(spid))

        unique_targets = sorted(list(set(target_ids)))
        return {
            "slide_number": slide_number,
            "has_timing": has_timing,
            "target_references_count": len(target_ids),
            "unique_animated_shapes_count": len(unique_targets),
            "target_shape_ids": unique_targets,
        }

    def _inspect_slide_relationships(self, slide) -> List[Dict[str, str]]:
        """Maps all relationships associated with the slide part."""
        rels_list = []
        try:
            for r_id, rel in slide.part.rels.items():
                rel_short = rel.reltype.split("/")[-1]
                rels_list.append({
                    "r_id": r_id,
                    "type": rel_short,
                    "target_ref": str(rel.target_ref)
                })
        except Exception as e:
            logger.warning("Could not read slide relationships: %s", e)
        return rels_list

    def _inspect_placeholders(self, slide, slide_number: int) -> List[Dict[str, Any]]:
        """Inspects placeholders and compares against slide layout placeholders."""
        results = []
        for shape in slide.shapes:
            if shape.is_placeholder:
                ph_xml = shape._element.xpath(".//p:ph")
                ph_type = ph_xml[0].get("type", "body") if ph_xml else "body"
                ph_idx = ph_xml[0].get("idx") if ph_xml else None

                # Find layout placeholder match
                layout_match = None
                try:
                    layout = slide.slide_layout
                    for l_shape in layout.shapes:
                        if l_shape.is_placeholder:
                            l_ph_xml = l_shape._element.xpath(".//p:ph")
                            l_type = l_ph_xml[0].get("type", "body") if l_ph_xml else "body"
                            l_idx = l_ph_xml[0].get("idx") if l_ph_xml else None
                            if (ph_idx and l_idx == ph_idx) or (ph_type and l_type == ph_type):
                                layout_match = f"LayoutShape_{l_shape.shape_id} (type={l_type}, idx={l_idx})"
                                break
                except Exception:
                    pass

                results.append({
                    "slide_number": slide_number,
                    "shape_id": shape.shape_id,
                    "name": shape.name,
                    "placeholder_type": ph_type,
                    "placeholder_index": ph_idx,
                    "text_preview": shape.text_frame.text[:60] if shape.has_text_frame else "",
                    "layout_placeholder_match": layout_match,
                    "position_pt": (emu_to_points(shape.left), emu_to_points(shape.top)),
                    "size_pt": (emu_to_points(shape.width), emu_to_points(shape.height)),
                })
        return results

    def _inspect_lists(self, slide, slide_number: int) -> List[Dict[str, Any]]:
        """Inspects paragraph OpenXML markup for bullet and numbering metadata."""
        list_items = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs, start=1):
                    p_elem = p._element
                    p_prs = p_elem.xpath(".//a:pPr")
                    if p_prs:
                        p_pr = p_prs[0]
                        bu_char = p_pr.xpath(".//a:buChar")
                        bu_autonum = p_pr.xpath(".//a:buAutoNum")
                        bu_none = p_pr.xpath(".//a:buNone")

                        bullet_type = "NONE"
                        bullet_char = None
                        autonum_type = None
                        autonum_start_at = None

                        if bu_none:
                            bullet_type = "EXPLICIT_NONE"
                        elif bu_char:
                            bullet_type = "CHAR"
                            bullet_char = bu_char[0].get("char")
                        elif bu_autonum:
                            bullet_type = "AUTONUM"
                            autonum_type = bu_autonum[0].get("type")
                            autonum_start_at = int(bu_autonum[0].get("startAt", "1"))
                        elif p.level > 0:
                            bullet_type = "INHERITED_LEVEL_BULLET"

                        mar_l = int(p_pr.get("marL", "0")) if p_pr.get("marL") else None
                        indent = int(p_pr.get("indent", "0")) if p_pr.get("indent") else None
                        lvl = int(p_pr.get("lvl", "0")) if p_pr.get("lvl") else p.level

                        if bullet_type != "NONE" or p.text.strip():
                            list_items.append({
                                "slide_number": slide_number,
                                "shape_id": shape.shape_id,
                                "paragraph_index": p_idx,
                                "text": p.text[:60],
                                "bullet_type": bullet_type,
                                "bullet_char": bullet_char,
                                "autonum_type": autonum_type,
                                "autonum_start_at": autonum_start_at,
                                "marL_emu": mar_l,
                                "indent_emu": indent,
                                "level": lvl,
                            })
        return list_items

    def _flatten_nodes(self, node: Dict[str, Any], result: List[Dict[str, Any]]) -> None:
        """Flattens a hierarchical shape tree into a list."""
        result.append(node)
        for child in node.get("children", []):
            self._flatten_nodes(child, result)

    def _collect_group_transforms(self, node: Dict[str, Any], result: List[Dict[str, Any]]) -> None:
        """Extracts group coordinate records for diagnostic reporting."""
        if node.get("is_group"):
            raw = node.get("raw_transform_data")
            if raw:
                children_info = []
                for child in node.get("children", []):
                    c_raw = child.get("raw_transform_data")
                    children_info.append({
                        "child_shape_id": child["shape_id"],
                        "child_name": child["name"],
                        "child_local_bbox_pt": child["local_bbox_pt"],
                        "child_computed_slide_bbox_pt": child["slide_bbox_pt"],
                    })
                result.append({
                    "group_shape_id": node["shape_id"],
                    "group_name": node["name"],
                    "off_emu": (raw["off_x_emu"], raw["off_y_emu"]),
                    "ext_emu": (raw["ext_cx_emu"], raw["ext_cy_emu"]),
                    "chOff_emu": (raw["chOff_x_emu"], raw["chOff_y_emu"]),
                    "chExt_emu": (raw["chExt_cx_emu"], raw["chExt_cy_emu"]),
                    "children": children_info,
                })
        for child in node.get("children", []):
            self._collect_group_transforms(child, result)

    def export_diagnostic_reports(self, inspection_data: Dict[str, Any], output_dir: Path) -> Tuple[Path, Path]:
        """Saves machine-readable JSON and human-readable Markdown diagnostic reports."""
        target_dir = Path(output_dir).resolve()
        target_dir.mkdir(parents=True, exist_ok=True)

        json_path = target_dir / "ooxml_inspection.json"
        md_path = target_dir / "ooxml_inspection.md"

        # 1. JSON Export
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(inspection_data, f, indent=2)

        # 2. Markdown Export
        md_lines = [
            f"# SlideGuard Deep OOXML Inspection Report",
            f"",
            f"**File Name:** `{inspection_data['file_name']}`  ",
            f"**Total Slides:** {inspection_data['total_slides']}  ",
            f"**Total Unique Shapes:** {inspection_data['global_id_inventory']['total_unique_shape_ids']}  ",
            f"**Total Animated Shapes:** {inspection_data['global_id_inventory']['total_animated_shape_ids']}  ",
            f"",
            f"---",
            f"",
            f"## 1. Global Preservation Constraints",
            f"",
            f"### Animated Shape IDs (MUST NOT BE DELETED OR RECREATED)",
            f"`{', '.join(str(i) for i in inspection_data['global_id_inventory']['animated_shape_ids'])}`",
            f"",
            f"---",
            f"",
            f"## 2. Per-Slide Shape Tree & Reading Order Hierarchy",
            f"",
        ]

        for s in inspection_data["slides"]:
            md_lines.extend([
                f"### Slide {s['slide_number']} (Layout: `{s['layout_name']}`)",
                f"- **Total Shapes:** {s['shape_count']} (Top-level in `p:spTree`: {s['top_level_element_count']})",
                f"- **Animated Shapes on Slide:** `{s['animation_info']['target_shape_ids']}`",
                f"",
                f"| Tree # | Tag | ID | Name | Type | Nesting | Computed Slide BBox (Pt) | Footer Status | Animated |",
                f"|---|---|---|---|---|---|---|---|---|",
            ])
            for shape in s["flat_shape_list"]:
                bbox = f"L:{shape['slide_bbox_pt']['left']}, T:{shape['slide_bbox_pt']['top']}, W:{shape['slide_bbox_pt']['width']}, H:{shape['slide_bbox_pt']['height']}"
                md_lines.append(
                    f"| {shape['tree_index']} | `{shape['xml_tag']}` | `{shape['shape_id']}` | {shape['name']} | {shape['object_type']} | Depth {shape['nesting_depth']} | {bbox} | `{shape['footer_candidate_type']}` | {'⚡ YES' if shape['is_animated'] else 'No'} |"
                )
            md_lines.append("")

            # Group Transforms section if any
            if s["group_transforms"]:
                md_lines.extend([
                    f"#### Group Transform Coordinate Calculations (Slide {s['slide_number']})",
                    f"",
                    f"| Group ID | Name | off (EMU) | ext (EMU) | chOff (EMU) | chExt (EMU) | Children |",
                    f"|---|---|---|---|---|---|---|",
                ])
                for grp in s["group_transforms"]:
                    md_lines.append(
                        f"| `{grp['group_shape_id']}` | {grp['group_name']} | `{grp['off_emu']}` | `{grp['ext_emu']}` | `{grp['chOff_emu']}` | `{grp['chExt_emu']}` | {len(grp['children'])} children |"
                    )
                md_lines.append("")

            # Lists section if any
            if s["list_items"]:
                md_lines.extend([
                    f"#### List Paragraph Structures (Slide {s['slide_number']})",
                    f"",
                    f"| Shape ID | Para # | Bullet Type | AutoNum Type | Start At | Level | Text Preview |",
                    f"|---|---|---|---|---|---|---|",
                ])
                for li in s["list_items"]:
                    md_lines.append(
                        f"| `{li['shape_id']}` | {li['paragraph_index']} | `{li['bullet_type']}` | `{li['autonum_type'] or 'N/A'}` | {li['autonum_start_at'] or 'N/A'} | {li['level']} | {li['text']} |"
                    )
                md_lines.append("")

        md_path.write_text("\n".join(md_lines), encoding="utf-8")
        return json_path, md_path
