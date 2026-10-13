"""
SlideGuard Presentation Reader.
Parses PowerPoint PPTX files into internal models with OOXML-level fidelity.
"""

from pathlib import Path
from typing import List, Optional, Tuple
import logging
import re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from models.enums import ShapeType, PlaceholderType, ListType
from models.shape import ShapeModel, TextParagraphModel, TextRunModel
from models.slide import SlideModel
from models.presentation import PresentationModel, DocumentMetadata
from engine.common.interfaces import IPresentationReader
from pptx_engine.utilities.geometry import emu_to_points, is_in_footer_region
from pptx_engine.ooxml.xml_utils import get_shape_alt_text, get_shape_decorative_status
from pptx_engine.ooxml.xml_namespaces import NAMESPACES

logger = logging.getLogger(__name__)


class PresentationReader(IPresentationReader):
    """Parses PPTX files safely without modifying source files."""

    def read(self, file_path: Path) -> PresentationModel:
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        try:
            prs = Presentation(str(path))
        except Exception as e:
            logger.error("Failed to load PowerPoint file %s: %s", path, e)
            raise ValueError(f"Corrupt or invalid PowerPoint file: {path} - {e}") from e

        slide_width_pt = emu_to_points(prs.slide_width)
        slide_height_pt = emu_to_points(prs.slide_height)

        # Extract core metadata
        metadata = self._extract_metadata(prs)

        model = PresentationModel(
            file_path=path,
            slide_width_pt=slide_width_pt,
            slide_height_pt=slide_height_pt,
            metadata=metadata,
            raw_presentation=prs
        )

        for idx, slide in enumerate(prs.slides, start=1):
            slide_model = self._parse_slide(slide, idx, slide_width_pt, slide_height_pt)
            model.slides.append(slide_model)

        return model

    def _extract_metadata(self, prs: Presentation) -> DocumentMetadata:
        meta = DocumentMetadata()
        try:
            core_props = prs.core_properties
            meta.title = core_props.title or None
            meta.author = core_props.author or None
            meta.subject = core_props.subject or None
            meta.keywords = core_props.keywords or None
            meta.comments = core_props.comments or None
            meta.created = str(core_props.created) if core_props.created else None
            meta.modified = str(core_props.modified) if core_props.modified else None
        except Exception as e:
            logger.warning("Could not parse core properties: %s", e)
        return meta

    def _parse_slide(
        self,
        slide,
        slide_number: int,
        slide_width_pt: float,
        slide_height_pt: float
    ) -> SlideModel:
        layout_name = "Unknown"
        try:
            if hasattr(slide, "slide_layout") and slide.slide_layout is not None:
                layout_name = getattr(slide.slide_layout, "name", "Layout")
        except Exception:
            pass

        slide_model = SlideModel(
            slide_id=slide.slide_id,
            slide_number=slide_number,
            layout_name=layout_name,
            raw_slide=slide
        )

        # Parse title if available
        try:
            title_shape = slide.shapes.title
            if title_shape is not None:
                slide_model.title_shape_id = title_shape.shape_id
                tx_body = None
                if hasattr(title_shape, "_element") and title_shape._element is not None:
                    tx_body = title_shape._element.find(f"{{{NAMESPACES['p']}}}txBody")
                if (tx_body is not None or not hasattr(title_shape, "_element")) and title_shape.has_text_frame:
                    slide_model.title = title_shape.text_frame.text.strip()
        except Exception:
            pass

        # Parse shapes
        for z_idx, shape in enumerate(slide.shapes):
            try:
                parsed_shape = self._parse_shape(shape, z_order=z_idx, slide_height_pt=slide_height_pt)
                slide_model.shapes.append(parsed_shape)

                # Check for footer candidacy
                if parsed_shape.is_placeholder and parsed_shape.placeholder_type == PlaceholderType.FOOTER:
                    slide_model.has_footer = True
                    slide_model.footer_text = parsed_shape.text_content
                    slide_model.footer_shape_ids.append(parsed_shape.shape_id)
                elif is_in_footer_region(parsed_shape.top_pt, parsed_shape.height_pt, slide_height_pt):
                    # Flag bottom text ONLY when combined with strong footer evidence (e.g. 'WILEY', copyright)
                    if parsed_shape.has_text and (
                        re.search(r"\b(wiley|copyright|©|\(c\)|all rights reserved)\b", parsed_shape.text_content, re.IGNORECASE)
                        or re.search(r"\b(footer|slide\s*num|page\s*num|ftr|sldnum)\b", parsed_shape.name, re.IGNORECASE)
                    ):
                        slide_model.has_footer = True
                        slide_model.footer_text = parsed_shape.text_content.strip()
                        slide_model.footer_shape_ids.append(parsed_shape.shape_id)

            except Exception as e:
                logger.warning("Error parsing shape on slide %d: %s", slide_number, e)

        return slide_model

    def _parse_shape(self, shape, z_order: int, slide_height_pt: float, parent_group_id: Optional[int] = None) -> ShapeModel:
        shape_type = ShapeType.OTHER
        placeholder_type = PlaceholderType.NONE
        is_placeholder = shape.is_placeholder

        if is_placeholder:
            shape_type = ShapeType.PLACEHOLDER
            try:
                ph_format = shape.placeholder_format
                ph_type = ph_format.type
                if ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                    placeholder_type = PlaceholderType.TITLE
                elif ph_type == PP_PLACEHOLDER.BODY:
                    placeholder_type = PlaceholderType.BODY
                elif ph_type == PP_PLACEHOLDER.SUBTITLE:
                    placeholder_type = PlaceholderType.SUBTITLE
                elif ph_type == PP_PLACEHOLDER.FOOTER:
                    placeholder_type = PlaceholderType.FOOTER
                elif ph_type == PP_PLACEHOLDER.HEADER:
                    placeholder_type = PlaceholderType.HEADER
                elif ph_type == PP_PLACEHOLDER.DATE:
                    placeholder_type = PlaceholderType.DATE_TIME
                elif ph_type == PP_PLACEHOLDER.SLIDE_NUMBER:
                    placeholder_type = PlaceholderType.SLIDE_NUMBER
                elif ph_type == PP_PLACEHOLDER.PICTURE:
                    placeholder_type = PlaceholderType.PICTURE
                elif ph_type == PP_PLACEHOLDER.TABLE:
                    placeholder_type = PlaceholderType.TABLE
                else:
                    placeholder_type = PlaceholderType.OBJECT
            except Exception:
                placeholder_type = PlaceholderType.OBJECT
        elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            shape_type = ShapeType.TEXT_BOX
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shape_type = ShapeType.PICTURE
        elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            shape_type = ShapeType.GROUP
        elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            shape_type = ShapeType.TABLE
        elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            shape_type = ShapeType.AUTO_SHAPE
        elif shape.shape_type == MSO_SHAPE_TYPE.MEDIA:
            shape_type = ShapeType.MEDIA

        # Geometry
        left_pt = emu_to_points(shape.left) if hasattr(shape, "left") else 0.0
        top_pt = emu_to_points(shape.top) if hasattr(shape, "top") else 0.0
        width_pt = emu_to_points(shape.width) if hasattr(shape, "width") else 0.0
        height_pt = emu_to_points(shape.height) if hasattr(shape, "height") else 0.0

        shape_model = ShapeModel(
            shape_id=shape.shape_id,
            name=shape.name or f"Shape_{shape.shape_id}",
            shape_type=shape_type,
            left_pt=left_pt,
            top_pt=top_pt,
            width_pt=width_pt,
            height_pt=height_pt,
            z_order=z_order,
            parent_group_id=parent_group_id,
            is_placeholder=is_placeholder,
            placeholder_type=placeholder_type,
            raw_element=shape._element if hasattr(shape, "_element") else None
        )

        # XML alt text and decorative status
        if hasattr(shape, "_element"):
            shape_model.alt_text = get_shape_alt_text(shape._element)
            shape_model.is_decorative = get_shape_decorative_status(shape._element)

        # Extract text & paragraphs
        if shape.has_text_frame:
            sp_elm = getattr(shape, "_element", None)
            tx_body = None
            if sp_elm is not None:
                tx_body = sp_elm.find(f"{{{NAMESPACES['p']}}}txBody")

            # Avoid calling shape.text_frame on shapes that have no txBody in XML (which would insert an empty txBody)
            if sp_elm is not None and tx_body is None:
                shape_model.text_content = ""
            else:
                tf = shape.text_frame
                full_text = []
                for p in tf.paragraphs:
                    p_text = p.text or ""
                    full_text.append(p_text)
                    p_lvl = self._get_paragraph_level(p)
                    p_model = TextParagraphModel(text=p_text, level=p_lvl)

                    # Identify list type
                    p_model.list_type = self._detect_list_type(p, level=p_lvl)

                    for r in p.runs:
                        font_name, font_size_pt, font_bold, font_italic, font_color_rgb = self._extract_run_formatting(r)
                        r_model = TextRunModel(
                            text=r.text or "",
                            font_name=font_name,
                            font_size_pt=font_size_pt,
                            font_bold=font_bold,
                            font_italic=font_italic,
                            font_color_rgb=font_color_rgb,
                        )
                        p_model.runs.append(r_model)
                    shape_model.paragraphs.append(p_model)
                shape_model.text_content = "\n".join(full_text)

        # Extract table info
        if shape.has_table:
            shape_model.table_row_count = len(shape.table.rows)
            shape_model.table_col_count = len(shape.table.columns)

        # Handle groups recursively
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP and hasattr(shape, "shapes"):
            for child_z, child_shape in enumerate(shape.shapes):
                child_model = self._parse_shape(
                    child_shape,
                    z_order=child_z,
                    slide_height_pt=slide_height_pt,
                    parent_group_id=shape.shape_id
                )
                shape_model.children.append(child_model)

        return shape_model

    def _detect_list_type(self, paragraph, level: Optional[int] = None) -> ListType:
        """Determines the list style of a paragraph from its formatting and content."""
        if level is None:
            level = self._get_paragraph_level(paragraph)
        text = (paragraph.text or "").strip()
        # Direct bullet check
        if level > 0:
            # Sub-levels in standard body placeholders default to bullets unless numbered
            if text and text[0].isdigit() and len(text) > 2 and text[1] in [".", ")"]:
                return ListType.NUMBERED_NUMERIC
            return ListType.BULLET

        # Check for numeric numbering pattern: "1. ", "1) "
        if text and len(text) > 2:
            first_word = text.split()[0]
            if first_word[:-1].isdigit() and first_word[-1] in [".", ")"]:
                return ListType.NUMBERED_NUMERIC
            if len(first_word) == 2 and first_word[0].isalpha() and first_word[1] in [".", ")"]:
                if first_word[0].islower():
                    return ListType.NUMBERED_ALPHA_LOWER
                else:
                    return ListType.NUMBERED_ALPHA_UPPER

        return ListType.NONE

    @staticmethod
    def _get_paragraph_level(paragraph) -> int:
        """Safely extracts paragraph level without mutating XML by adding <a:pPr/>."""
        p_elm = getattr(paragraph, "_p", None)
        if p_elm is None:
            p_elm = getattr(paragraph, "_element", None)
        if p_elm is not None:
            pPr = p_elm.find(f"{{{NAMESPACES['a']}}}pPr")
            if pPr is not None:
                lvl = pPr.get("lvl")
                if lvl is not None:
                    try:
                        return int(lvl)
                    except (ValueError, TypeError):
                        return 0
            return 0
        try:
            return getattr(paragraph, "level", 0) or 0
        except Exception:
            return 0

    @staticmethod
    def _extract_run_formatting(run) -> Tuple[Optional[str], Optional[float], Optional[bool], Optional[bool], Optional[str]]:
        """Safely extracts run character properties without mutating XML by adding <a:rPr/> or <a:solidFill/>.

        Returns:
            (font_name, font_size_pt, font_bold, font_italic, font_color_rgb)
        """
        r_elm = getattr(run, "_r", None)
        if r_elm is None:
            r_elm = getattr(run, "_element", None)

        if r_elm is not None:
            rPr = r_elm.find(f"{{{NAMESPACES['a']}}}rPr")
            if rPr is None:
                return None, None, None, None, None

            font_name = None
            latin = rPr.find(f"{{{NAMESPACES['a']}}}latin")
            if latin is not None:
                font_name = latin.get("typeface") or None

            font_size_pt = None
            sz = rPr.get("sz")
            if sz is not None:
                try:
                    font_size_pt = float(sz) / 100.0
                except (ValueError, TypeError):
                    pass

            font_bold = None
            b = rPr.get("b")
            if b is not None:
                font_bold = b in ("1", "true", "True")

            font_italic = None
            i = rPr.get("i")
            if i is not None:
                font_italic = i in ("1", "true", "True")

            font_color_rgb = None
            solid_fill = rPr.find(f"{{{NAMESPACES['a']}}}solidFill")
            if solid_fill is not None:
                srgb_clr = solid_fill.find(f"{{{NAMESPACES['a']}}}srgbClr")
                if srgb_clr is not None:
                    val = srgb_clr.get("val")
                    if val:
                        font_color_rgb = val.upper()

            return font_name, font_size_pt, font_bold, font_italic, font_color_rgb

        # Fallback for mock objects without underlying XML
        font_name = None
        font_size_pt = None
        font_bold = None
        font_italic = None
        font_color_rgb = None
        if hasattr(run, "font") and run.font:
            font_name = getattr(run.font, "name", None)
            sz = getattr(run.font, "size", None)
            if sz is not None:
                try:
                    font_size_pt = emu_to_points(sz)
                except Exception:
                    pass
            font_bold = getattr(run.font, "bold", None)
            font_italic = getattr(run.font, "italic", None)
            try:
                color = getattr(run.font, "color", None)
                if color and hasattr(color, "rgb") and color.rgb:
                    font_color_rgb = str(color.rgb)
            except Exception:
                pass

        return font_name, font_size_pt, font_bold, font_italic, font_color_rgb
